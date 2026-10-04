#!/usr/bin/env python3
"""싱크에이전트 「Sync Agent 설치 사전 질의서」(엑셀)를 만든다. SoT = PREINSTALL-QUESTIONNAIRE.md §1 (여기와 1:1).

고객사가 칸을 채워 돌려보내는 양식이라 PDF 가 아니라 엑셀이다(2026-10-01 Harold 지시 · 게이트웨이
tools/generate_agent_environment_checklist_xlsx.py 와 같은 틀). 회신 구성 그대로 리허설 픽스처를 만들어
그 고객사 조합을 빌드 · 검증한 뒤에만 출고한다(FEATURE-SYNC-AGENT §8). 그래서 고객 · 구매 표 구조를 통째로 받는다.

고객 전달물이므로 우리 쪽 주소 · 포트 · 서버 이름 · 내부 경로 · 다른 고객사 이름은 싣지 않는다. 줄표(—)도 쓰지 않는다.

사용법:
    python build-questionnaire-xlsx.py [판]
"""

import hashlib
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
DEFAULT_VERSION = "1.1"
TITLE = "Sync Agent 설치 사전 질의서"

FONT_NAME = "맑은 고딕"
CODE_FONT = "Consolas"
GREEN = "047857"
INPUT_FILL = PatternFill("solid", fgColor="FFF6D6")
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
SECTION_FILL = PatternFill("solid", fgColor="DDF3EA")
LABEL_FILL = PatternFill("solid", fgColor="FAFBFC")
NOTE_FILL = PatternFill("solid", fgColor="FFF6E5")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center")

# (번호, 항목, 고르는 값 목록 또는 None, 작성 예 · 확인 방법) — PREINSTALL-QUESTIONNAIRE.md §1 과 1:1
SECTIONS = [
    ("1. 에이전트를 설치할 서버", [
        ("1-1", "운영체제와 정확한 버전", None, "예: Windows Server 2016 Datacenter x64\n「확인 명령」 시트 1번"),
        ("1-2", "운영체제 언어", ["한국어판", "영문판", "기타"], "「확인 명령」 시트 1번(코드 페이지 949 = 한국어판)"),
        ("1-3", "이 서버에서 인터넷(HTTPS, 443 포트) 접속", ["가능", "방화벽 신청 필요"], ""),
        ("1-4", "관리자 권한으로 설치 진행", ["가능", "불가(별도 협의 필요)"], ""),
    ]),
    ("2. 연동할 데이터베이스", [
        ("2-1", "DB 종류와 정확한 버전", None,
         "예: MySQL 8.0 / Aurora MySQL 3.x / SQL Server 2016 / Oracle 11g / PostgreSQL 15\n「확인 명령」 시트 2번"),
        ("2-2", "클라우드 DB 여부", ["사내 서버", "AWS RDS · Aurora", "Azure", "기타 클라우드"], "기타 클라우드면 서비스 이름을 적어 주세요"),
        ("2-3", "암호화(TLS) 연결 강제 여부", ["예", "아니오", "모름"], "「확인 명령」 시트 2번"),
        ("2-4", "에이전트 설치 서버에서 DB 호스트 · 포트로 접속", ["가능", "방화벽 신청 필요"], ""),
    ]),
    ("3. 고객 · 구매 데이터 구성 (가장 중요합니다)", [
        ("3-1", "고객(회원) 데이터의 위치: DB(스키마)명과 테이블 또는 뷰 이름", None, "예: shop.member_view (뷰)"),
        ("3-2", "구매(주문) 데이터의 위치: DB(스키마)명과 테이블 또는 뷰 이름", None, "예: shop.order_view (뷰)"),
        ("3-3", "원본이 접속 DB와 다른 DB에 있어 뷰로 모아 노출하는 구성", ["예", "아니오"], "예: 고객 · 주문 원본이 각각 다른 DB"),
        ("3-4", "대략 몇 행입니까", None, "예: 고객 약 30만 · 주문 약 500만\n「확인 명령」 시트 4번"),
        ("3-5", "한 행을 고유하게 식별하는 컬럼", None,
         "예: 고객 = 회원번호 / 주문 = 주문번호 + 항목순번\n뷰로 연동하시면 이 컬럼이 뷰에 반드시 있어야 변경분이 자동 반영됩니다"),
        ("3-6", "등록 · 수정 일시 컬럼의 이름과, 값에 시각(시:분)까지 있는지", None,
         "예: 고객 upd_dt (시각 포함) / 구매 판매일 (날짜만)"),
        ("3-7", "회원이 탈퇴하면 고객 데이터에서 어떻게 처리됩니까", ["행을 지운다", "탈퇴 표시 컬럼에 남긴다", "모름"],
         "표시 컬럼에 남긴다면 컬럼 이름과 값을 오른쪽 답 칸에 함께 적어 주세요"),
        ("3-8", "문자(광고) 수신동의 컬럼의 이름과 실제로 들어 있는 값 전부", None,
         "예: SMS_YN (Y · N · 공란)\n「확인 명령」 시트 5번 결과를 그대로 적어 주세요"),
        ("3-9", "고객 · 구매 표(뷰)의 구조", ["시트에 붙였음", "파일로 첨부함"],
         "「확인 명령」 시트 3번 결과를 「고객 표 구조」 · 「구매 표 구조」 시트에 붙여 주세요"),
    ]),
    ("4. 접속 계정 (비밀번호는 적지 마세요)", [
        ("4-1", "에이전트용 읽기 전용 계정 이름 (새로 만드는 것을 권합니다)", None, "예: sync_reader"),
        ("4-2", "그 계정에 SELECT 권한을 줄 대상", ["3-1 · 3-2 의 테이블(뷰) 2개", "기타"], "기타면 대상을 적어 주세요"),
        ("4-3", "계정의 접속 허용 호스트", ["에이전트 서버 IP 지정", "내부망 전체"], ""),
    ]),
    ("5. 담당자 · 일정", [
        ("5-1", "설치 진행 담당자 성함 · 연락처 · 이메일", None, ""),
        ("5-2", "희망 설치 일정", None, "예: 10월 셋째 주 평일 19시 이후"),
    ]),
]

# (구분, 확인할 것, 명령) — 고객 · 구매 표 이름 자리에 3-1 · 3-2 의 실제 이름을 넣는다
COMMAND_GROUPS = [
    ("1. 서버 (Windows · 명령 프롬프트)", [
        ("운영체제 이름과 버전", 'systeminfo | findstr /B /C:"OS"'),
        ("32비트 · 64비트", "echo %PROCESSOR_ARCHITECTURE%"),
        ("운영체제 언어(949 = 한국어판)", "chcp"),
    ]),
    ("1. 서버 (리눅스)", [
        ("운영체제 이름과 버전", "cat /etc/*release"),
        ("32비트 · 64비트", "uname -m"),
    ]),
    ("2. 데이터베이스 버전과 암호화 연결 (데이터베이스에 접속해서 실행)", [
        ("MySQL · Aurora 버전", "SELECT VERSION();"),
        ("MySQL · Aurora 암호화 연결 강제", "SHOW VARIABLES LIKE 'require_secure_transport';"),
        ("SQL Server 버전", "SELECT @@VERSION;"),
        ("Oracle 버전", "SELECT banner FROM v$version;"),
        ("PostgreSQL 버전", "SELECT version();"),
        ("PostgreSQL 암호화 연결", "SHOW ssl;"),
    ]),
    ("3. 표 구조 (고객 표와 구매 표 둘 다 · 표이름 자리에 실제 이름)", [
        ("MySQL · Aurora (테이블 · 뷰 공통)", "SHOW CREATE TABLE 표이름;"),
        ("SQL Server", "EXEC sp_help '표이름';"),
        ("Oracle", "DESC 표이름"),
        ("PostgreSQL", "\\d 표이름"),
    ]),
    ("4. 행 수 (고객 표와 구매 표 둘 다)", [
        ("모든 DB 공통", "SELECT COUNT(*) FROM 표이름;"),
    ]),
    ("5. 수신동의 컬럼에 실제로 들어 있는 값", [
        ("모든 DB 공통 (컬럼이름 · 고객표이름 자리에 실제 이름)", "SELECT 컬럼이름, COUNT(*) FROM 고객표이름 GROUP BY 컬럼이름;"),
    ]),
]


def font(size=10, bold=False, color="1A1A1A", name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, color=color)


def put(ws, ref, value, *, fnt=None, fill=None, align=WRAP_TOP, border=True):
    cell = ws[ref]
    cell.value = value
    cell.font = fnt or font()
    cell.alignment = align
    if fill:
        cell.fill = fill
    if border:
        cell.border = BORDER
    return cell


def header(ws, row, labels):
    for col, label in zip("ABCDEFGH", labels):
        put(ws, f"{col}{row}", label, fnt=font(bold=True, color="FFFFFF"), fill=HEAD_FILL, align=WRAP_CENTER)
    ws.row_dimensions[row].height = 22


def title_block(ws, title, subtitle, last_col):
    ws.merge_cells(f"A1:{last_col}1")
    put(ws, "A1", title, fnt=font(16, True, GREEN), align=Alignment(vertical="center"), border=False)
    ws.row_dimensions[1].height = 32
    ws.merge_cells(f"A2:{last_col}2")
    put(ws, "A2", subtitle, fnt=font(9, color="5C6673"), align=Alignment(vertical="center"), border=False)


def fit_page(ws, orientation):
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_form(ws, version):
    ws.title = "작성"
    for col, width in zip("ABCD", (7, 40, 44, 54)):
        ws.column_dimensions[col].width = width
    title_block(ws, TITLE, f"한줄로 데이터 동기화 에이전트 · v{version}", "D")

    notes = [
        "설치 전에 아래 내용을 회신해 주시면, 귀사 환경과 같은 구성으로 저희가 먼저 설치 전 과정을 검증한 뒤 설치 파일을 전달드립니다.",
        "노란 칸만 적어 주세요. 고르는 칸은 칸을 누르면 목록이 나옵니다. 목록에 없으면 직접 적으셔도 됩니다.",
        "모르는 항목은 비워 두셔도 됩니다. 담당자가 함께 확인해 드립니다.",
        "비밀번호 · 접속 암호는 적지 마세요.",
    ]
    row = 3
    for index, text in enumerate(notes):
        ws.merge_cells(f"A{row}:D{row}")
        last = index == len(notes) - 1
        put(ws, f"A{row}", text, fnt=font(9.5, bold=last, color="5C3D00" if last else "1A1A1A"),
            fill=NOTE_FILL if last else None, align=WRAP_CENTER, border=False)
        ws.row_dimensions[row].height = 20
        row += 1
    row += 1

    header_row = row
    header(ws, row, ("번호", "항목", "답 (노란 칸에 적어 주세요)", "작성 예 · 확인 방법"))
    ws.freeze_panes = f"A{row + 1}"
    row += 1

    for section_title, items in SECTIONS:
        ws.merge_cells(f"A{row}:D{row}")
        put(ws, f"A{row}", section_title, fnt=font(10.5, True, "064E3B"), fill=SECTION_FILL, align=WRAP_CENTER)
        for col in "BCD":
            ws[f"{col}{row}"].border = BORDER
        ws.row_dimensions[row].height = 22
        row += 1
        for number, item, options, example in items:
            put(ws, f"A{row}", number, fill=LABEL_FILL)
            put(ws, f"B{row}", item, fill=LABEL_FILL)
            put(ws, f"C{row}", None, fill=INPUT_FILL)
            put(ws, f"D{row}", example, fnt=font(9, color="5C6673"))
            if options:
                # 목록 밖의 값(기타 · 컬럼 이름 덧붙임)도 적을 수 있어야 하므로 오류 창은 띄우지 않는다.
                validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                                            allow_blank=True, showErrorMessage=False)
                ws.add_data_validation(validation)
                validation.add(f"C{row}")
            # 작성 예 · 항목이 칸 높이에 잘리지 않게 줄 수만큼 높인다(한 줄에 한글 약 34자 · 게이트웨이 양식 실측).
            def lines(text, width):
                return sum(max(1, -(-len(part) // width)) for part in text.split("\n")) if text else 1
            ws.row_dimensions[row].height = max(26, 15 * max(lines(example, 34), lines(item, 25)) + 6)
            row += 1

    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    put(ws, f"A{row}",
        "다 적으신 뒤 이 파일을 담당자에게 메일로 보내 주세요. 회신 주신 구성 그대로 저희 쪽에서 설치 전 과정을 검증한 뒤 설치 파일을 전달드립니다.",
        fnt=font(9.5), align=WRAP_CENTER, border=False)
    ws.row_dimensions[row].height = 30

    fit_page(ws, "portrait")
    ws.print_title_rows = f"{header_row}:{header_row}"


def build_commands(ws):
    for col, width in zip("ABC", (7, 44, 80)):
        ws.column_dimensions[col].width = width
    title_block(ws, "확인 명령", "아래 명령을 실행한 뒤, 화면에 나온 글자를 「작성」 시트의 해당 칸이나 구조 시트에 그대로 붙여 주세요.", "C")
    ws.merge_cells("A3:C3")
    put(ws, "A3", "명령이 동작하지 않으면 그 항목은 비워 두셔도 됩니다. 담당자가 함께 확인해 드립니다.",
        fnt=font(9.5), align=WRAP_CENTER, border=False)
    row = 5
    header(ws, row, ("", "확인할 것", "명령"))
    row += 1
    for group_title, commands in COMMAND_GROUPS:
        ws.merge_cells(f"A{row}:C{row}")
        put(ws, f"A{row}", group_title, fnt=font(10.5, True, "064E3B"), fill=SECTION_FILL, align=WRAP_CENTER)
        for col in "BC":
            ws[f"{col}{row}"].border = BORDER
        ws.row_dimensions[row].height = 22
        row += 1
        for label, command in commands:
            put(ws, f"A{row}", None, fill=LABEL_FILL)
            put(ws, f"B{row}", label, fill=LABEL_FILL)
            # 한글이 섞인 명령은 고정폭 글꼴에 한글이 없으므로 본문 글꼴로 둔다.
            ascii_only = all(ord(ch) < 0x80 for ch in command)
            cell = put(ws, f"C{row}", command, fnt=font(10, name=CODE_FONT) if ascii_only else font())
            cell.number_format = "@"
            ws.row_dimensions[row].height = 20
            row += 1
    fit_page(ws, "landscape")


def build_paste(ws, which):
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 22
    title_block(ws, f"{which} 구조", f"「확인 명령」 시트 3번을 {which}에 실행한 결과를 아래 노란 칸(A5부터)에 그대로 붙여 주세요.", "H")
    ws.merge_cells("A3:H3")
    put(ws, "A3", "결과가 길면 파일로 저장해 메일에 첨부하셔도 됩니다. 데이터(행 값)는 붙이지 않으셔도 됩니다.",
        fnt=font(9.5), align=WRAP_CENTER, border=False)
    for row in range(5, 65):
        for col in "ABCDEFGH":
            cell = ws[f"{col}{row}"]
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.font = font(9.5)
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top")
    fit_page(ws, "landscape")


def build(version, output):
    workbook = Workbook()
    build_form(workbook.active, version)
    build_commands(workbook.create_sheet("확인 명령"))
    build_paste(workbook.create_sheet("고객 표 구조"), "고객 표")
    build_paste(workbook.create_sheet("구매 표 구조"), "구매 표")
    workbook.properties.title = f"{TITLE} v{version}"
    workbook.properties.creator = "INVITO"
    workbook.properties.lastModifiedBy = "INVITO"
    workbook.save(output)


def main():
    version = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_VERSION
    output = ROOT / f"SyncAgent_사전질의서_v{version.replace('.', '_')}.xlsx"
    build(version, output)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    print(f"SYNC_AGENT_QUESTIONNAIRE_XLSX_OK {output.name} bytes={output.stat().st_size} sha256={digest[:12]}")


if __name__ == "__main__":
    main()
