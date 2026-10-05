# -*- coding: utf-8 -*-
"""
재접수 제출 문서 변환기 (2026-10-03) — 제출문서/*.md → .docx

하는 일
  1. 문서 맨 위 「작성 메모」 블록(<!-- … -->)을 걷어 낸다(내부 메모는 제출본에 싣지 않는다).
  2. pandoc 으로 docx 를 만들고, python-docx 로 글꼴(맑은 고딕) · 표 테두리 · 표 머리 음영을 맞춘다.
  3. 본문 검사 — 줄표 · 「예정」 · 「시범」 · 「확인 필요」 · 코드 식별자 · 빈 표 칸이 남아 있으면 그 문서 이름을 경고로 찍는다(변환은 한다).

실행(저장소 루트): python docs/transmission-cert/resubmit-1006/build_submission_docx.py [출력 폴더]
기본 출력: C:\\Users\\ceo\\Downloads\\전송자격인증제\\제출본_docx_<오늘 날짜>
⛔ 메모에 「확정 대기」가 남은 문서는 미리보기다. 결정 · 배포가 끝난 뒤 다시 돌려 최종본을 만든다.
"""
import datetime
import os
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '제출문서')
FONT = '맑은 고딕'

CHECKS = {
    '줄표': r'—',
    '예정': r'예정',
    '시범': r'시범',
    '확인 필요': r'확인 필요',
    '코드 식별자': r'`|\.ts\b|\.tsx\b|routes/|utils/',
    '빈 표 칸': r'\|\s+\|\s+\|',
}


def strip_memo(text):
    if text.lstrip().startswith('<!--'):
        end = text.find('-->')
        if end >= 0:
            return text[end + 3:].lstrip('\n'), True
    return text, False


def body_warnings(body):
    out = []
    for label, pat in CHECKS.items():
        n = sum(1 for line in body.splitlines() if re.search(pat, line))
        if n:
            out.append(f'{label} {n}줄')
    return out


def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '808080')
        borders.append(el)
    tbl_pr.append(borders)


def set_run_font(run, size=None, bold=None):
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        r_pr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(attr), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def polish(path):
    doc = Document(path)
    for style in doc.styles:
        try:
            if style.font is not None:
                style.font.name = FONT
                r_pr = style.element.get_or_add_rPr()
                fonts = r_pr.find(qn('w:rFonts'))
                if fonts is None:
                    fonts = OxmlElement('w:rFonts')
                    r_pr.append(fonts)
                for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                    fonts.set(qn(attr), FONT)
        except Exception:
            pass
    for p in doc.paragraphs:
        for r in p.runs:
            set_run_font(r)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                if ri == 0:
                    set_cell_shading(cell, 'E8ECF2')
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r, size=9, bold=True if ri == 0 else None)
    doc.save(path)


def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.expanduser('~'), 'Downloads', '전송자격인증제',
        '제출본_docx_' + datetime.date.today().strftime('%Y%m%d'))
    os.makedirs(out_dir, exist_ok=True)
    for name in sorted(os.listdir(SRC)):
        if not name.endswith('.md'):
            continue
        text = open(os.path.join(SRC, name), encoding='utf-8').read()
        body, had_memo = strip_memo(text)
        warns = body_warnings(body)
        with tempfile.NamedTemporaryFile('w', suffix='.md', delete=False, encoding='utf-8') as tmp:
            tmp.write(body)
            tmp_path = tmp.name
        target = os.path.join(out_dir, name[:-3] + '.docx')
        subprocess.run(['pandoc', tmp_path, '-f', 'gfm', '-t', 'docx', '--resource-path', SRC, '-o', target], check=True)
        os.unlink(tmp_path)
        polish(target)
        flag = (' · 경고: ' + ', '.join(warns)) if warns else ''
        memo = ' · 작성 메모 있었음(확정 대기 확인)' if had_memo else ''
        print(f'{os.path.basename(target)}{memo}{flag}')
    print('출력:', out_dir)


if __name__ == '__main__':
    main()
