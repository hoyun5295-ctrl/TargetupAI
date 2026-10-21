# -*- coding: utf-8 -*-
"""전송자격인증 현장 실사(2026-10-22) 로그 명령집 엑셀 생성기.

엑셀 칸 하나에 여러 줄을 넣고 복사하면 앞뒤에 " 가 붙고 안의 " 가 "" 로 바뀐다(2026-10-21 이 PC 엑셀 실측).
그래서 명령은 전부 한 줄이다. SQL 은 bash here-string(<<< "...") 으로 넘긴다.
  - SQL 안에 " $ ` ! 를 쓰지 않는다(바깥 큰따옴표가 깨지거나 셸이 바꾼다) → 별칭은 따옴표 없는 한글_밑줄.
  - 한국 시각 = docker exec -e PGTZ=Asia/Seoul (SET 줄을 따로 찍지 않게).
조회는 전부 SELECT 다. 실행: python make_inspection_commands_xlsx.py
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / '실사_로그_명령집_20261022.xlsx'

HANJUL_PSQL = 'docker exec -i -e PGTZ=Asia/Seoul targetup-postgres psql -U targetup -d targetup -P pager=off -P border=2'
GW_PSQL = 'docker exec -i -e PGTZ=Asia/Seoul bito-bench-postgres psql -U bito -d bito_gateway -P pager=off -P border=2'
# 보여 주는 범위 = 한줄로 · 비토 엔진만(★2026-10-21 Harold). 외주 중계로 나가는 발송 DB 라인 1~11 은 넣지 않는다 → 발송 이력 = G-11.

H_ADMIN = '한줄로 서버(.62) · administrator'
H_ROOT = '한줄로 서버(.62) · root 셸'
G_INVITO = '게이트웨이 서버(.65) · invito'
G_ROOT = '게이트웨이 서버(.65) · root 셸'

OLD = '10-03 실행분'
NEW = '새로 씀 · 10-21 컬럼 확인'
NEW_RAW = '새로 씀'


def one_line(sql: str) -> str:
    s = re.sub(r'\s+', ' ', sql).strip()
    for ch in '"$`!':
        assert ch not in s, f'SQL 안에 {ch} 가 있다: {s[:80]}'
    return s


def pg(prefix: str, sql: str) -> str:
    return f'{prefix} <<< "{one_line(sql)}"'


ROWS = []  # (번호, 항목, 로그, 위치, 명령, 근거, 메모)


def add(no, item, title, where, cmd, basis, memo=''):
    assert '\n' not in cmd, f'{no} 명령이 여러 줄이다'
    ROWS.append((no, item, title, where, cmd, basis, memo))


# ───────────────────────────── 한줄로 ─────────────────────────────
add('H-1', '4.1 · 4.2', '감사 기록 보관 범위 · 월별 건수', H_ADMIN, pg(HANJUL_PSQL, """
SELECT MIN(created_at) AS 가장_오래된_기록, MAX(created_at) AS 가장_최근_기록, COUNT(*) AS 전체_건수 FROM audit_logs;
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, COUNT(*) AS 건수 FROM audit_logs GROUP BY 1 ORDER BY 1;
"""), OLD, '지우는 기능이 없다는 근거 = 가장 오래된 기록이 서비스 초기(2026-02-11)부터 남아 있음')

add('H-2', '4.1 · 4.2', '기록 종류별 첫 기록일 · 월별 건수', H_ADMIN, pg(HANJUL_PSQL, """
SELECT action AS 기록_종류, to_char(MIN(created_at), 'YYYY-MM-DD') AS 첫_기록일,
       COUNT(*) FILTER (WHERE created_at >= '2026-08-01' AND created_at < '2026-09-01') AS 건수_8월,
       COUNT(*) FILTER (WHERE created_at >= '2026-09-01' AND created_at < '2026-10-01') AS 건수_9월,
       COUNT(*) FILTER (WHERE created_at >= '2026-10-01') AS 건수_10월
  FROM audit_logs GROUP BY action ORDER BY action;
"""), OLD, '어떤 로그가 언제부터 남는지 한 장으로 보여 줄 때')

add('H-3', '4.1', '로그인 성공 · 실패 · 차단 · 동시 접속 원문 + 로그인 차단 대장', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(u.login_id, s.login_id, a.details->>'loginId') AS 계정,
       host(a.ip_address) AS 출발지_IP, left(a.user_agent, 40) AS 브라우저
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('login_success','login_fail','login_blocked','logout','login_session_conflict','login_takeover')
 ORDER BY a.created_at DESC LIMIT 30;
SELECT b.blocked_at AS 차단_시각, b.login_id AS 계정, host(b.ip_address) AS 출발지_IP, b.reason AS 사유, b.fail_count AS 실패_수,
       b.expires_at AS 만료, b.unblocked_at AS 해제, b.unblock_reason AS 해제_사유,
       COALESCE(s2.login_id, u2.login_id) AS 해제한_사람
  FROM login_blocks b LEFT JOIN super_admins s2 ON s2.id = b.unblocked_by LEFT JOIN users u2 ON u2.id = b.unblocked_by
 ORDER BY b.blocked_at DESC LIMIT 20;
"""), '10-21 실행분 + 해제한 사람 칸 새로 씀', '로그인 차단 = 같은 IP · 아이디 10분 안 5회 실패 시 30분 자동 차단 · 30분 경과 또는 관리자 해제')

add('H-4', '4.1', '화면(페이지) 접속 기록', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, COUNT(*) AS 화면_접속_기록, COUNT(DISTINCT user_id) AS 계정_수
  FROM audit_logs WHERE action = 'page_view' GROUP BY 1 ORDER BY 1;
SELECT a.created_at AS 일시, COALESCE(u.login_id, s.login_id) AS 계정, a.details->>'path' AS 화면, host(a.ip_address) AS 출발지_IP
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'page_view' ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분', '2026-09-12부터 기록 · 같은 사람 · 같은 화면 60초 안은 1건 · 10-21 실측 = 9월 7,758(135계정) · 10월 8,122(145계정)')

add('H-5', '4.2', '개인정보 조회 · 등록 · 수정 · 내보내기 · 삭제 기록', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, action AS 종류, details->>'kind' AS 대상, COUNT(*) AS 요청_수
  FROM audit_logs
 WHERE action IN ('privacy_view','privacy_edit','privacy_export','privacy_purge','customer_delete','customer_bulk_delete','customer_delete_all','customer_delete_by_user')
 GROUP BY 1, 2, 3 ORDER BY 1, 2, 3;
SELECT a.created_at AS 일시, a.action AS 종류, a.details->>'kind' AS 대상, a.details->>'count' AS 건수,
       COALESCE(u.login_id, s.login_id) AS 계정, host(a.ip_address) AS 출발지_IP
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('privacy_view','privacy_edit','privacy_export','privacy_purge','customer_delete_all')
 ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분', '조회 · 수정은 2026-09-12부터 · 내보내기는 2026-08-20부터 · 삭제는 2026-02부터 · 0건 조회도 기록')

add('H-6', '2.1 · 4.1 ②', '계정 발급(회원가입) · 고객사 등록 기록 + 월별 계정 생성 수', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(s.login_id, u.login_id) AS 발급자, a.details AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('user_account_created', 'company_created') ORDER BY a.created_at;
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, COUNT(*) AS 생성된_계정
  FROM users WHERE COALESCE(is_system, false) = false GROUP BY 1 ORDER BY 1;
"""), OLD, '가입이 아니라 법인 계약 뒤 당사 발급 · 발급 기록은 2026-10-02부터, 그 전 계정은 생성일')

add('H-6b', '2.1', '월별 · 계정 유형별 발급 수와 고객사 수(4월 · 7월 몰림 설명)', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(created_at, 'YYYY-MM') AS 월, user_type AS 계정_유형, COUNT(*) AS 계정_수, COUNT(DISTINCT company_id) AS 고객사_수
  FROM users WHERE COALESCE(is_system, false) = false GROUP BY 1, 2 ORDER BY 1, 2;
"""), '10-21 실행분(4 · 7월만 본 것을 전체 달로)', '관리자 수 = 고객사 수 → 회사 등록과 함께 관리자 1개씩 발급')

add('H-6c', '2.1', '지워진 첫 시연 회사(10/09 17:28 등록)를 가리키는 감사 기록 전부', H_ADMIN, pg(HANJUL_PSQL, """
SELECT created_at AS 일시, action AS 종류, target_type AS 대상, left(details::text, 120) AS 상세 FROM audit_logs
 WHERE target_id::text = 'eb33d30f-a494-4a24-bb84-7a639fb522ee' OR details::text LIKE '%eb33d30f-a494-4a24-bb84-7a639fb522ee%'
 ORDER BY created_at;
"""), '10-21 실행분', '등록 · 화면 접속 15줄뿐 · 삭제 줄 없음(정리 스크립트가 감사 기록을 안 남긴다) → 물어보면 답 Q-06')

add('H-6d', '2.1', '첫 시연 회사 삭제 때 남긴 백업 파일', H_ADMIN,
    'ls -l ~/demo-purge-eb33d30f-a494-4a24-bb84-7a639fb522ee-*.jsonl', '10-21 실행분(administrator · root 둘 다 파일 없음)',
    '없다고 나오면 root 셸에서 같은 명령(스크립트를 root 로 돌렸으면 그쪽 홈) · 파일 시각 = 삭제 시각')

add('H-7', '2.1 · 3.4 ②', '본인확인기관(한국모바일인증) 인증 이력', H_ADMIN, pg(HANJUL_PSQL, """
SELECT v.created_at AS 시작, u.login_id AS 계정, v.purpose AS 구분, v.status AS 상태, v.provider AS 인증기관,
       left(v.verified_name, 1) || '**' AS 인증_이름,
       left(v.verified_phone, 3) || '****' || right(v.verified_phone, 4) AS 인증_번호,
       left(v.provider_tx_id, 8) AS 요청번호_앞8, (v.dup_key_hash IS NOT NULL) AS 중복확인값_기록,
       v.verified_at AS 완료, v.ip_address AS 접속_IP
  FROM identity_verifications v LEFT JOIN users u ON u.id = v.user_id
 ORDER BY v.created_at DESC LIMIT 30;
"""), '10-21 실행분 + 중복확인값 칸 새로 씀', '이름 첫 글자 · 번호 가운데 · 요청번호 앞 8자리만 나온다 · 운영 개시 2026-10-06 · 중복확인값은 있다/없다만(값은 안 나온다)')

add('H-7b', '2.1 ② · 3.4 ②', '본인확인 계정의 지금 수신 번호 = 본인확인 번호인가(번호는 안 찍는다)', H_ADMIN, pg(HANJUL_PSQL, """
SELECT u.login_id AS 계정, (regexp_replace(COALESCE(u.mfa_phone, ''), '[^0-9]', '', 'g') = v.verified_phone) AS 수신번호_본인확인번호_일치
  FROM users u JOIN LATERAL (SELECT verified_phone FROM identity_verifications WHERE user_id = u.id AND status = 'verified'
                              ORDER BY verified_at DESC LIMIT 1) v ON true
 ORDER BY 1;
"""), '10-21 실행분', '10-21 실측 = 5계정 전부 t · 물어보면 답 Q-61')

add('H-8', '2.1', '발신번호 등록 신청 · 심사 / 계정별 배정 / 회선 정책 변경', H_ADMIN, pg(HANJUL_PSQL, """
SELECT r.created_at AS 신청, c.company_name AS 고객사, u.login_id AS 신청자, r.phone AS 발신번호, r.number_type AS 번호_유형,
       jsonb_array_length(CASE WHEN jsonb_typeof(r.documents) = 'array' THEN r.documents ELSE '[]'::jsonb END) AS 서류_수,
       r.status AS 결과, s.login_id AS 심사자, r.reviewed_at AS 심사_일시, r.reject_reason AS 반려_사유
  FROM sender_registrations r LEFT JOIN companies c ON c.id = r.company_id
  LEFT JOIN users u ON u.id = r.requested_by LEFT JOIN super_admins s ON s.id = r.reviewed_by
 ORDER BY r.created_at DESC LIMIT 30;
SELECT a.created_at AS 배정_일시, c.company_name AS 고객사, n.phone AS 발신번호, u.login_id AS 대상_계정,
       COALESCE(s.login_id, b.login_id) AS 배정한_관리자
  FROM callback_number_assignments a JOIN callback_numbers n ON n.id = a.callback_number_id
  LEFT JOIN companies c ON c.id = n.company_id LEFT JOIN users u ON u.id = a.user_id
  LEFT JOIN super_admins s ON s.id = a.assigned_by LEFT JOIN users b ON b.id = a.assigned_by
 ORDER BY a.created_at DESC LIMIT 30;
SELECT a.created_at AS 일시, COALESCE(s.login_id, u.login_id) AS 처리자, a.details AS 회선_정책_변경_내용
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'sender_line_policy_update' ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분 + 서류 수 · 번호 유형 칸 새로 씀', '발신번호 등록 · 수정 · 삭제 · 배정 변경 이력은 10-21 부터 H-8c · 물어보면 답 Q-26')

add('H-8b', '2.1', '시험 번호 02-1234-5678 · 두 번 승인된 052-937-6647 이 지금 발신번호로 몇 개 등록돼 있나', H_ADMIN, pg(HANJUL_PSQL, """
SELECT c.company_name AS 고객사, n.phone AS 발신번호, n.created_at AT TIME ZONE 'UTC' AS 등록
  FROM callback_numbers n LEFT JOIN companies c ON c.id = n.company_id
 WHERE regexp_replace(n.phone, '[^0-9]', '', 'g') IN ('0212345678', '0529376647')
 ORDER BY 2, 3;
"""), '10-21 실행분(한국 시각 표시까지 확인)', '물어보면 답 Q-23 · Q-24 · 발신번호 등록 시각은 UTC 저장이라 한국 시각으로 바꿔 보인다(10-21 · 8/19 04:07 UTC = 13:07 KST · 심사 시각과 소수점까지 같음)')

add('H-8c', '2.1 ⑤ · 4.2', '발신번호 등록 · 수정 · 기본 번호 · 삭제 · 배정 변경 이력(10-21 부터)', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, s.login_id AS 처리자, c.company_name AS 고객사,
       a.details->>'phone' AS 번호, a.details->>'via' AS 경로, a.details->>'isDefault' AS 기본번호
  FROM audit_logs a LEFT JOIN super_admins s ON s.id = a.user_id
  LEFT JOIN companies c ON c.id::text = a.details->>'companyId'
 WHERE a.action LIKE 'callback_%' ORDER BY a.created_at DESC LIMIT 30;
"""), '10-21 실행분(배포 뒤 실측 · 시간 조건만 뺌)', '10-21 13:15:37 등록 · 13:15:45 삭제(ceo · 테스트계정 · 02-1234-5678 · 경로 admin) = S1 실측 · 그 전 기록은 없다(Q-26) · 고객사 화면에서 한 일은 처리자가 고객 계정이라 처리자 칸이 비고 경로 = company')

add('H-9', '2.2 ① · ④', '국외 IP 차단 — 국내 대역 · 날짜별 감지와 차단 · 차단 원문', H_ADMIN, pg(HANJUL_PSQL, """
SELECT country_code AS 국가, source AS 출처, COUNT(*) AS 대역_수, MAX(updated_at) AS 마지막_갱신
  FROM geo_allow_cidrs GROUP BY 1, 2 ORDER BY 3 DESC;
SELECT to_char(created_at, 'YYYY-MM-DD') AS 날짜,
       COUNT(*) FILTER (WHERE action = 'foreign_access_detected') AS 감지_예외통과_포함,
       COUNT(*) FILTER (WHERE action = 'foreign_access_blocked') AS 차단
  FROM audit_logs WHERE action IN ('foreign_access_detected', 'foreign_access_blocked') AND created_at >= '2026-09-01'
 GROUP BY 1 ORDER BY 1;
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(u.login_id, s.login_id, a.details->>'loginId') AS 계정,
       host(a.ip_address) AS 출발지_IP, a.details AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'foreign_access_blocked' ORDER BY a.created_at DESC LIMIT 30;
"""), OLD, '탐지 2026-08-28 · 차단 2026-10-01 · ⚠ 10/07 15:57~16:06 국내로 보이는 출발지 차단 · 그 직후 16:07 대역 갱신 → H-9b · Q-32')

add('H-9b', '2.2', '10/07 차단 주소의 지금 국내 대역 여부 · 국내 대역 교체 이력 · 10/01 뒤 감지가 예외 통과인지', H_ADMIN, pg(HANJUL_PSQL, """
SELECT v.ip AS 출발지_IP, EXISTS (SELECT 1 FROM geo_allow_cidrs g WHERE v.ip::inet <<= g.cidr) AS 지금_국내_대역
  FROM (VALUES ('14.52.61.97'), ('61.81.163.98'), ('115.138.27.202'), ('153.72.39.239'), ('219.164.64.138')) AS v(ip);
SELECT a.created_at AS 일시, COALESCE(s.login_id, u.login_id) AS 처리자, left(a.details::text, 150) AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'geo_cidrs_replaced' ORDER BY a.created_at;
SELECT to_char(created_at, 'YYYY-MM-DD') AS 날짜, details->>'exempted' AS 예외_통과, COUNT(*) AS 건수
  FROM audit_logs WHERE action = 'foreign_access_detected' AND created_at >= '2026-10-01' GROUP BY 1, 2 ORDER BY 1, 2;
"""), '10-21 실행분', '물어보면 답 Q-32 · Q-33 · 10/07 15:47 교체 2597 → 1 기록 · 16:07 되돌림은 기록 없음')

add('H-10', '2.2 ③', '허용 IP(예외) 대장 — 사유 · 승인자 · 허용 기간 + 허용 · 해제 이력', H_ADMIN, pg(HANJUL_PSQL, """
SELECT x.created_at AS 등록_일시, x.scope AS 범위, c.company_name AS 고객사, u.login_id AS 계정, x.cidr AS 허용_대역,
       x.reason AS 사유, s.login_id AS 승인자, x.approved_at AS 승인_일시, x.expires_at AS 허용_만료, x.is_active AS 유효
  FROM access_origin_allowlist x LEFT JOIN companies c ON c.id = x.company_id
  LEFT JOIN users u ON u.id = x.user_id LEFT JOIN super_admins s ON s.id = x.approved_by
 ORDER BY x.created_at;
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(s.login_id, u.login_id) AS 처리자, a.details AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('access_origin_exception_granted', 'access_origin_exception_revoked', 'geo_cidrs_replaced')
 ORDER BY a.created_at;
"""), OLD, '회수는 지우지 않고 유효=false 로 남긴다')

add('H-11', '2.2 ⑤', '연동 경로(API · 싱크에이전트) 출발지 대조 기록 — 월별', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, action AS 종류, details->>'scope' AS 경로,
       details->>'registered' AS 등록_여부, COUNT(*) AS 건수, COUNT(DISTINCT target_id) AS 고객사_수,
       COUNT(DISTINCT ip_address) AS 출발지_수
  FROM audit_logs WHERE action IN ('machine_origin_detected', 'machine_origin_blocked')
 GROUP BY 1, 2, 3, 4 ORDER BY 1, 2, 3;
"""), OLD, '사람은 국가로, 기계(API · 싱크)는 등록한 출발지로 통제 · 10-21 실측 = 회사 API 미등록 0 · 싱크에이전트 10월 미등록 60건(2곳 · 출발지 2) · 같은 출발지 1시간 1건 기록')

add('H-11b', '2.2 ⑤', '10월 미등록 출발지 — 회사 · 경로 · 주소 · 처음 · 마지막', H_ADMIN, pg(HANJUL_PSQL, """
SELECT c.company_name AS 고객사, a.details->>'scope' AS 경로, host(a.ip_address) AS 출발지_IP, COUNT(*) AS 건수,
       MIN(a.created_at) AS 처음, MAX(a.created_at) AS 마지막
  FROM audit_logs a LEFT JOIN companies c ON c.id::text = a.target_id::text
 WHERE a.action = 'machine_origin_detected' AND a.details->>'registered' = 'false' AND a.created_at >= '2026-10-01'
 GROUP BY 1, 2, 3 ORDER BY 6 DESC;
"""), '10-21 실행분', '결함 D-16 · 10-21 실측 = 이새 10/01~10/03 10:00(등록 전) · 테스트계정 115.138.27.202 10/15(당사 시험) · 10/15 23:13 뒤 0')

add('H-12', '2.3', '입금 명의 보류 · 고객사 소명 · 처리 결과', H_ADMIN, pg(HANJUL_PSQL, """
SELECT d.created_at AS 신청, c.company_name AS 고객사, d.amount AS 금액, left(d.depositor_name, 1) || '**' AS 입금자,
       d.held_reason AS 보류_사유, d.held_at AS 보류_시각, d.explanation_note AS 고객사_소명, d.status AS 결과,
       d.confirmed_at AS 처리_일시, d.admin_note AS 관리자_메모
  FROM deposit_requests d LEFT JOIN companies c ON c.id = d.company_id
 WHERE d.held_at IS NOT NULL OR d.held_reason IS NOT NULL ORDER BY d.created_at DESC LIMIT 20;
SELECT a.created_at AS 일시, COALESCE(s.login_id, u.login_id) AS 처리자, a.details->>'resolution' AS 결과,
       left(a.details->>'depositorName', 1) || '**' AS 입금자, a.details->>'heldReason' AS 보류_사유,
       a.details->>'explanationNote' AS 고객사_소명, a.details->>'adminNote' AS 관리자_메모
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'deposit_hold_resolved' ORDER BY a.created_at;
"""), '10-21 실행분 + 처리 기록 칸별로 새로 씀', '입금자 이름은 두 표 모두 첫 글자만(10-21 처리 기록 상세가 이름 원문을 찍어 고침 = D-18)')

add('H-13', '3.4 ① · ③', '다중인증 요청 · 성공 · 실패 · 잠김 + 원문 + 수신 번호 변경 · 본인인증 기록', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월,
       COUNT(*) FILTER (WHERE action = 'mfa_challenge') AS 인증_요청, COUNT(*) FILTER (WHERE action = 'mfa_success') AS 인증_성공,
       COUNT(*) FILTER (WHERE action = 'mfa_fail') AS 인증_실패, COUNT(*) FILTER (WHERE action = 'mfa_locked') AS 잠김,
       COUNT(DISTINCT user_id) FILTER (WHERE action = 'mfa_success') AS 인증한_계정_수
  FROM audit_logs WHERE action IN ('mfa_challenge', 'mfa_success', 'mfa_fail', 'mfa_locked') GROUP BY 1 ORDER BY 1;
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(u.login_id, a.details->>'loginId') AS 계정, host(a.ip_address) AS 출발지_IP
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id
 WHERE a.action IN ('mfa_challenge', 'mfa_success', 'mfa_fail', 'mfa_locked') ORDER BY a.created_at DESC LIMIT 20;
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(s.login_id, u.login_id) AS 처리자, t.login_id AS 대상_계정
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
  LEFT JOIN users t ON t.id = a.target_id
 WHERE a.action IN ('mfa_phone_changed', 'identity_verify_start', 'identity_verified', 'identity_verify_fail')
 ORDER BY a.created_at DESC LIMIT 40;
"""), '10-21 실행분', '5회 실패 = 잠김 · 인증번호는 평문 저장 안 함 · 10-21 실측 = 9월 요청 56 · 성공 52 · 실패 1 · 잠김 0 / 10월 114 · 109 · 0 · 0 · 6계정')

add('H-13b', '3.4 ③', '10월 인증번호 수신 번호 변경 — 대상 · 변경 전후(저장 때 가림)', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, COALESCE(s.login_id, u.login_id) AS 처리자, a.target_type AS 대상_종류, a.target_id AS 대상_id,
       left(a.details::text, 200) AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action = 'mfa_phone_changed' AND a.created_at >= '2026-10-01' ORDER BY a.created_at;
"""), '10-21 실행분', '10-21 실측 = 10/13 gwchae 전후 같음 · 10/03 invito01 · 10/20 efolium(옛 계정) 전후 비어 있음 · 10/20 acehardware01 비어 있음 → 등록')

add('H-14', '3.5 ①', '발신 인증(추가 인증) 요청 · 성공 + 원문', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월,
       COUNT(*) FILTER (WHERE action = 'sender_auth_challenge') AS 인증_요청,
       COUNT(*) FILTER (WHERE action = 'sender_auth_success') AS 인증_성공,
       COUNT(DISTINCT user_id) FILTER (WHERE action = 'sender_auth_success') AS 인증한_계정_수
  FROM audit_logs WHERE action IN ('sender_auth_challenge', 'sender_auth_success') GROUP BY 1 ORDER BY 1;
SELECT a.created_at AS 일시, a.action AS 종류, u.login_id AS 계정, host(a.ip_address) AS 출발지_IP, left(a.details::text, 150) AS 상세
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id
 WHERE a.action IN ('sender_auth_challenge', 'sender_auth_success') ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분', '인증 유지 8시간 · 문서 3.5 6절 「발신번호 · 수신 번호(가림)가 남는다」를 상세 칸으로 보인다')

add('H-15', '3.2 · 3.3 · 4.2', '직원 계정 목록 · 생성 · 등급 변경 · 사용 중지 · 비밀번호 변경 이력', H_ADMIN, pg(HANJUL_PSQL, """
SELECT login_id AS 직원_계정, name AS 이름, role AS 등급, is_active AS 사용, totp_enabled AS OTP_등록,
       must_change_password AS 비밀번호_변경_필요, created_at AT TIME ZONE 'UTC' AS 생성,
       last_login_at AT TIME ZONE 'UTC' AS 마지막_로그인
  FROM super_admins ORDER BY created_at;
SELECT a.created_at AS 일시, a.action AS 종류, s.login_id AS 처리자, a.details AS 상세
  FROM audit_logs a LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('admin_account_created','admin_account_disabled','admin_account_enabled','admin_role_changed','admin_password_changed','totp_enrolled')
 ORDER BY a.created_at DESC LIMIT 40;
"""), '10-21 실행분(한국 시각 표시까지 확인)', '등급 = super 대표 · lead 지원팀장 · support 지원팀원 · 직원 계정 표 시각은 UTC 저장이라 한국 시각으로 바꿔 보인다(10-21 · suran 생성 소수점까지 감사 기록과 같음) · 비밀번호 변경 기록에 비밀번호 값은 없다(코드 확인)')

add('H-16', '3.1 ④ · 3.3 ③ · 4.3', '운영 기록 대장 — 로그 점검 · 방화벽 정책 변경 · 접근권한 점검', H_ADMIN, pg(HANJUL_PSQL, """
SELECT c.created_at AS 작성, c.details->>'kind' AS 종류, c.details->>'system' AS 대상_시스템,
       left(c.details->>'title', 40) AS 제목, c.details->'recorder'->>'name' AS 작성자,
       f.created_at AS 확인_일시, fs.name AS 확인자, v.created_at AS 무효_처리
  FROM audit_logs c
  LEFT JOIN audit_logs f ON f.action = 'ops_record_confirmed' AND f.target_type = 'ops_record' AND f.target_id = c.target_id
  LEFT JOIN super_admins fs ON fs.id = f.user_id
  LEFT JOIN audit_logs v ON v.action = 'ops_record_voided' AND v.target_type = 'ops_record' AND v.target_id = c.target_id
 WHERE c.action = 'ops_record_created' AND c.target_type = 'ops_record' AND c.created_at >= '2026-10-01'
 ORDER BY c.created_at;
"""), '10-21 실행분', '화면(슈퍼관리자 → 보안 · 인증 → 운영 기록 대장)이 더 읽기 쉽다. 이건 지울 수 없는 감사 기록으로 쌓인다는 원본. 종류 log_review 로그 점검 · firewall_change 방화벽 변경 · access_review 접근권한 점검')

add('H-17', '4.3', '월간 로그 점검 자료(점검 때 보는 집계) + 로그인 실패 몰린 계정', H_ADMIN, pg(HANJUL_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 점검_대상_월,
       COUNT(*) FILTER (WHERE action = 'login_success') AS 로그인_성공, COUNT(*) FILTER (WHERE action = 'login_fail') AS 로그인_실패,
       COUNT(*) FILTER (WHERE action = 'login_blocked') AS 로그인_차단, COUNT(*) FILTER (WHERE action = 'login_session_conflict') AS 동시_접속_감지,
       COUNT(*) FILTER (WHERE action = 'mfa_locked') AS 인증_잠김, COUNT(*) FILTER (WHERE action = 'foreign_access_blocked') AS 국외_차단,
       COUNT(*) FILTER (WHERE action = 'machine_origin_blocked') AS 기계_접속_차단, COUNT(*) FILTER (WHERE action = 'account_restricted') AS 계정_제한,
       COUNT(*) FILTER (WHERE action = 'company_terminated') AS 회사_해지, COUNT(*) FILTER (WHERE action = 'privacy_export') AS 개인정보_내보내기,
       COUNT(*) FILTER (WHERE action = 'customer_delete_all') AS 고객_전체삭제, COUNT(*) FILTER (WHERE action = 'admin_role_changed') AS 직원_등급_변경,
       COUNT(*) FILTER (WHERE action = 'admin_account_created') AS 직원_계정_생성
  FROM audit_logs WHERE created_at >= '2026-09-01' GROUP BY 1 ORDER BY 1;
SELECT * FROM (
  SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, details->>'loginId' AS 계정, COUNT(*) AS 실패_수,
         COUNT(DISTINCT ip_address) AS 출발지_수,
         ROW_NUMBER() OVER (PARTITION BY date_trunc('month', created_at) ORDER BY COUNT(*) DESC) AS rn
    FROM audit_logs WHERE action = 'login_fail' AND created_at >= '2026-09-01'
   GROUP BY date_trunc('month', created_at), details->>'loginId' HAVING COUNT(*) >= 5
) t WHERE rn <= 10 ORDER BY 월, 실패_수 DESC;
"""), OLD, '')

add('H-18', '5.1', '계정 제한 · 정지 · 고객사 해지 조치와 고지 결과', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, COALESCE(s.login_id, u.login_id) AS 처리자, a.details AS 사유_잠근계정_끊은세션_고지
  FROM audit_logs a LEFT JOIN users u ON u.id = a.user_id LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('account_restricted', 'company_terminated', 'kakao_profile_disabled')
 ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분', '10-21 실측 = 고객사 해지 2건(9/07 · 10/06)에 사유 · 고지 결과 없음 = D-25 · 계정 제한 1건(10/03 당사 시험 · notified true) · 10-21 배포 뒤 해지는 사유 · 안내 문자가 칸별로 남는다 = H-18b')

add('H-18b', '5.1 · 3.3', '고객사 해지 기록 · 해지 사유 · 안내 문자 결과(10-21 부터)', H_ADMIN, pg(HANJUL_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, s.login_id AS 처리자, a.details->>'companyName' AS 고객사,
       a.details->>'reason' AS 사유, a.details->>'note' AS 메모, a.details->>'before' AS 변경전,
       a.details->>'lockedUsers' AS 잠근_계정, a.details->>'killedSessions' AS 끊은_접속,
       a.details->>'notified' AS 문자_보냄, a.details->>'notifyFailed' AS 문자_실패
  FROM audit_logs a LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('company_terminated', 'company_termination_notice')
 ORDER BY a.created_at DESC LIMIT 20;
"""), '10-21 실행분(배포 뒤 실측 · 시간 조건만 뺌)', '10-21 13:39:31 테스트주식회사 해지 실측 = 사유 admin_action · 변경전 active · 잠근 계정 2 · 안내 문자 1 · 실패 0 · 대표 휴대폰 수신 · 사유 = contract_end 계약 종료 · spam_report 불법스팸 신고 · admin_action 관리자 조치 · 옛 해지 2건은 사유 · 문자 칸이 없다')

add('H-19', '5.2', '금칙어 규칙 · 월별 탐지 · 차단 · 원문 · 규칙 변경 이력', H_ADMIN, pg(HANJUL_PSQL, """
SELECT name AS 규칙, mode AS 모드, jsonb_array_length(elements) AS 조합_요소_수, is_active AS 사용, created_at AS 등록, updated_at AS 마지막_수정
  FROM spam_block_rules ORDER BY created_at;
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, mode AS 규칙_모드, action_taken AS 처리,
       COUNT(*) AS 발송_요청_수, SUM(affected_rows) AS 걸린_건수
  FROM spam_block_hits GROUP BY 1, 2, 3 ORDER BY 1, 2, 3;
SELECT h.created_at AS 일시, r.name AS 규칙, c.company_name AS 고객사, h.send_source AS 발송_경로, h.action_taken AS 처리,
       h.affected_rows AS 건수, h.content_sample AS 문안_표본_숫자가림
  FROM spam_block_hits h LEFT JOIN spam_block_rules r ON r.id = h.rule_id LEFT JOIN companies c ON c.id = h.company_id
 ORDER BY h.created_at DESC LIMIT 20;
SELECT a.created_at AS 일시, a.action AS 종류, s.login_id AS 처리자, a.details AS 상세
  FROM audit_logs a LEFT JOIN super_admins s ON s.id = a.user_id
 WHERE a.action IN ('spam_block_rule_create', 'spam_block_rule_update', 'spam_block_rule_mode') ORDER BY a.created_at;
"""), '10-21 실행분', '모드 detect = 탐지(기록만) · block = 차단 · 키워드 · URL · 전화번호 2~5개 조합 · 10-21 13:53 불법스팸 block 전환(D-26 닫힘) · 시험 규칙은 detect · 기록 1건 = 10/03 시험 차단')

add('H-21a', '3.1 ①', '방화벽 정책 · 서비스 가동 · 정책 파일 수정 시각', H_ROOT,
    "ufw status verbose; systemctl is-active ufw fail2ban; stat -c '%y  %n' /etc/ufw/user.rules /etc/ufw/after.rules", OLD)
add('H-21b', '3.1 ④', '방화벽 변경 명령 기록(서버 인증 로그)', H_ROOT,
    "zgrep -h -a 'COMMAND=.*ufw ' /var/log/auth.log* 2>/dev/null | sort | tail -60", OLD,
    '변경 내역 전체는 운영 기록 대장(H-16)의 방화벽 정책 변경')
add('H-21c', '3.1 ④', '서버 관리자 접속 이력(최근 30건)', H_ROOT, 'last -n 30 -F -i | head -34', OLD)
add('H-21d', '3.1 ①', '외부로 열린 포트', H_ROOT, "ss -tlnH | awk '{print $4}' | sort -u", OLD)
add('H-21e', '3.1 ① · 4.1 ⑤', 'DB 컨테이너가 서버 안(127.0.0.1)에만 열림', H_ADMIN,
    "docker ps --format 'table {{.Names}}\\t{{.Status}}\\t{{.Ports}}'", '10-21 실행분',
    '10-21 실측 = postgres · mysql · redis 127.0.0.1 · ⚠ pay-ingest-db 0.0.0.0:23388(결제 통계 수신 전용 · 컨테이너 방화벽 체인 = H-21h · D-29)')
add('H-21h', '3.1 ① · 4.1 ⑤', '결제 통계 수신 DB 포트의 컨테이너 전용 방화벽 체인(허용 출발지)', H_ROOT,
    'iptables -S DOCKER-USER', '10-21 실행분',
    '문서 3.1 4절 「송신 서버 5곳만 허용 · 그 밖은 차단」의 증거 · 10-21 실측 = 23388 허용 5곳(139.150.81.213 · .65 · .58 · .57 · .54) 뒤 DROP')
add('H-21g', '3.1 ③', '사설 원격 접속망(Tailscale)이 꺼져 있는가', H_ADMIN,
    'systemctl is-active tailscaled; systemctl is-enabled tailscaled', '10-21 실행분(점검 때 active · enabled → 10-21 두 서버 끔 · 대장 4건 = D-19 닫힘)',
    '꺼져 있으면 inactive · disabled(설치 안 됐으면 찾을 수 없다고 나온다) · 문서 3.1 6절 「별도 원격 접속 경로를 두지 않는다」의 증거')
add('H-21f', '3.1 ①', '침입 차단(SSH 반복 실패 차단) 현황', H_ROOT, 'fail2ban-client status sshd', '10-21 실행분(누적 실패 3 · 차단 0)',
    '감시 대상 이름이 sshd 가 아니면 fail2ban-client status 로 목록부터')

add('H-22a', '4.1 ⑤ · 4.2', '앱 로그 회전 설정 · 가장 오래된 파일 · 개수 · 용량 · 폴더 권한', H_ADMIN,
    "pm2 conf pm2-logrotate | grep -E 'retain|compress|max_size|rotateInterval'; ls -ltr ~/.pm2/logs | sed -n '2,6p'; "
    "echo \"파일 수: $(ls ~/.pm2/logs | wc -l) / 용량: $(du -sh ~/.pm2/logs | cut -f1)\"; ls -ld ~/.pm2/logs", '10-21 실행분',
    '400일 · 매일 회전 · 압축 · 10-21 실측 = 9/12 이후 179개 · 335M · 폴더 drwxr-xr-x')
add('H-22b', '4.1 ⑤ · 4.2', '웹 서버 로그 · 시스템 저널 1년 보관 설정과 가장 오래된 기록', H_ROOT,
    "grep -E '^\\s*(daily|weekly|rotate|compress)' /etc/logrotate.d/nginx; ls -tr /var/log/nginx | head -3; "
    "grep -E '^(SystemMaxUse|MaxRetentionSec)' /etc/systemd/journald.conf; journalctl --disk-usage; journalctl --no-pager -o short-iso | head -1",
    '10-21 실행분', '웹 서버 400회분 · 저널 1년 · 12G · 설정일 2026-09-12 · 10-21 파일 개수 한도 2,000 으로 고침 · 그 전 개수 한도(100)로 08-15 이전분 삭제 = D-31 · 대장 기록 · 물어보면 Q-84')
add('H-22c', '4.2', 'DB 서버 로그(최근 24시간 · 마지막 30줄)', H_ADMIN,
    'docker logs --since 24h targetup-postgres 2>&1 | tail -30', '10-21 실행분',
    '10-21 배포 뒤 = is_active 오류 0 · imc_template_status 는 재기동 직후와 6시간마다 1건만(기억 · 기능 영향 없음) · 그 전 기록에는 몇 분마다 찍혀 있다(D-28 · Q-80)')
add('H-22d', '4.2', 'DB 컨테이너 로그 설정', H_ADMIN,
    "docker inspect targetup-postgres --format '{{.HostConfig.LogConfig.Type}} {{json .HostConfig.LogConfig.Config}}'", '10-21 실행분(json-file {})',
    '설정이 비어 있으면(제한 없음) 잘리지 않고 남는다')

add('H-23', '4.1 ⑤', '백업 예약 · 실행 기록 · 서버 보관분 · 마지막 성공 · 감시 기록 · 경보 설정 여부', H_ADMIN,
    "crontab -l | grep -i backup; tail -12 /home/administrator/backups/cron.log; ls -lt /home/administrator/backups | sed -n '2,9p'; "
    "stat -c '%y  LAST_SUCCESS' /home/administrator/backups/LAST_SUCCESS; tail -8 /home/administrator/backups/monitor.log; "
    "grep -c '^ALERT_CMD=.\\+' /home/administrator/backups/.env", '10-21 실행분',
    '마지막 숫자 1 = 경보 보낼 곳 설정됨(값은 안 찍는다) · 10-21 실측 = 03:00 시작 · 03:13 성공 · PG 414M · MySQL 447M 암호화 · .59 전송 · 감시 매일 OK · ⚠ mysqldump 비밀번호 인자 경고 = D-30')

add('H-24', '2.2 · 3.4 · 3.5', '통제 시행 스위치(운영 설정 파일의 해당 줄만)', H_ADMIN,
    "grep -E '^(GEO_BLOCK_ENFORCE_FROM|ORIGIN_ALLOWLIST_ENFORCE_FROM|MFA_ENFORCE_FROM|MFA_PILOT_LOGIN_IDS|SENDER_AUTH_ENFORCE_FROM|"
    "SENDER_AUTH_PILOT_LOGIN_IDS|SENDER_AUTH_FAIL_CLOSED|IDENTITY_VERIFY_ENFORCE_FROM|IDENTITY_VERIFY_PILOT_LOGIN_IDS|IDENTITY_VERIFY_PROVIDER)=' "
    "/home/administrator/targetup-app/packages/backend/.env", OLD,
    '비밀번호 · 키 줄은 고르지 않는다. 값은 시행일과 계정 명단뿐 · ⛔ 실사 중에 cat .env 는 절대 치지 않는다')

# ───────────────────────────── 게이트웨이 ─────────────────────────────
add('G-1', '4.1 · 4.2', '감사 기록 보관 범위 · 월별 건수', G_INVITO, pg(GW_PSQL, """
SELECT MIN(created_at) AS 가장_오래된_기록, MAX(created_at) AS 가장_최근_기록, COUNT(*) AS 전체_건수 FROM audit_log;
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, COUNT(*) AS 건수 FROM audit_log GROUP BY 1 ORDER BY 1;
"""), OLD, '2026-06-15부터 · 지우는 코드 없음')

add('G-2', '4.1 · 4.2', '기록 종류별 첫 기록일 · 마지막 기록일 · 월별 건수', G_INVITO, pg(GW_PSQL, """
SELECT action AS 기록_종류, target_type AS 대상, to_char(MIN(created_at), 'YYYY-MM-DD') AS 첫_기록일,
       to_char(MAX(created_at), 'YYYY-MM-DD') AS 마지막_기록일,
       COUNT(*) FILTER (WHERE created_at >= '2026-08-01' AND created_at < '2026-09-01') AS 건수_8월,
       COUNT(*) FILTER (WHERE created_at >= '2026-09-01' AND created_at < '2026-10-01') AS 건수_9월,
       COUNT(*) FILTER (WHERE created_at >= '2026-10-01') AS 건수_10월
  FROM audit_log GROUP BY action, target_type ORDER BY action, target_type;
"""), OLD)

add('G-3', '2.2 ① · ③', '접속 계정별 허용 IP 지정 현황 · 변경 사유 · 허용 만료일', G_INVITO, pg(GW_PSQL, """
SELECT COUNT(*) AS 전체_계정, COUNT(*) FILTER (WHERE is_active) AS 사용_중,
       COUNT(*) FILTER (WHERE is_active AND COALESCE(array_length(allowed_ips, 1), 0) > 0) AS 사용중_허용IP_지정,
       COUNT(*) FILTER (WHERE is_active AND COALESCE(array_length(allowed_ips, 1), 0) = 0) AS 사용중_허용IP_없음
  FROM agent_account;
SELECT agent_id AS 계정, is_active AS 사용, COALESCE(array_length(allowed_ips, 1), 0) AS 허용_IP_수,
       allowed_ips_reason AS 변경_사유, allowed_ips_expires_at AS 허용_만료일, updated_at AS 마지막_수정
  FROM agent_account ORDER BY is_active DESC, agent_id;
"""), '10-21 실행분', 'IP 자체는 안 찍고 개수만 · 사유 · 만료일 칸은 2026-10-03 배포분 · 오늘 리허설 R-2 로 컬럼 확인')

add('G-4', '2.2 ① · ⑤ · 4.1', '접속 이력 — 월별 · 창구별 · 거부 사유 · 거부 원문', G_INVITO, pg(GW_PSQL, """
SELECT to_char(date_trunc('month', first_seen_at), 'YYYY-MM') AS 월, channel AS 창구, event_type AS 사건, outcome AS 결과,
       SUM(event_count) AS 건수, COUNT(DISTINCT remote_ip) AS 출발지_수
  FROM account_access_event GROUP BY 1, 2, 3, 4 ORDER BY 1, 2, 3, 4;
SELECT reason AS 거부_사유, channel AS 창구, SUM(event_count) AS 건수, COUNT(DISTINCT remote_ip) AS 출발지_수,
       MIN(first_seen_at) AS 처음, MAX(last_seen_at) AS 마지막
  FROM account_access_event WHERE outcome = 'denied' GROUP BY 1, 2 ORDER BY 3 DESC;
SELECT last_seen_at AS 일시, channel AS 창구, agent_id AS 계정, remote_ip AS 출발지_IP, reason AS 거부_사유, event_count AS 건수
  FROM account_access_event WHERE outcome = 'denied' ORDER BY last_seen_at DESC LIMIT 20;
"""), OLD, '10월 초 거부 급증 = 사내 시험 서버 옛 에이전트(물어보면 답 시트)')

add('G-4b', '2.2 ⑤ · 4.3', 'API 연동 10/03 뒤 사유별(미등록 출발지 기록 여부) · 9월 거부 계정 · 주소', G_INVITO, pg(GW_PSQL, """
SELECT channel AS 창구, outcome AS 결과, COALESCE(reason, '없음') AS 사유, SUM(event_count) AS 건수,
       COUNT(DISTINCT remote_ip) AS 출발지_수, MAX(last_seen_at) AS 마지막
  FROM account_access_event WHERE channel = 'api' AND first_seen_at >= '2026-10-03' GROUP BY 1, 2, 3 ORDER BY 1, 2, 3;
SELECT agent_id AS 계정, remote_ip AS 출발지_IP, reason AS 거부_사유, SUM(event_count) AS 건수,
       MIN(first_seen_at) AS 처음, MAX(last_seen_at) AS 마지막
  FROM account_access_event WHERE outcome = 'denied' AND first_seen_at >= '2026-09-01' AND first_seen_at < '2026-10-01'
 GROUP BY 1, 2, 3 ORDER BY 4 DESC;
"""), '10-21 실행분', '결함 D-16 · Q-42 · 10-21 실측 = 10/03 뒤 API 7,944건 전부 사유 없음 · 9월 거부 5줄')

add('G-5', '2.2 ③', '접속 계정 변경 이력(허용 IP 등록 · 변경 · 만료)', G_INVITO, pg(GW_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, u.username AS 처리자, a.target_id AS 대상_계정,
       (a.detail ? 'allowed_ips') AS 허용IP_변경_포함, a.detail->'allowed_ips' AS 바꾼_허용IP
  FROM audit_log a LEFT JOIN admin_user u ON u.id = a.admin_user_id
 WHERE a.target_type = 'agent_account' AND a.action IN ('UPDATE','ACTIVATE','DEACTIVATE','SET_RESTRICTIONS','HARD_DELETE','ALLOWED_IPS_EXPIRED')
 ORDER BY a.created_at DESC LIMIT 40;
"""), OLD)

add('G-6', '3.1 ④', '서버 접속 · 명령 수집 기록 + 관리 사이트 로그인 기록', G_INVITO, pg(GW_PSQL, """
SELECT action AS 종류, to_char(MIN(created_at), 'YYYY-MM-DD HH24:MI') AS 첫_기록, to_char(MAX(created_at), 'YYYY-MM-DD HH24:MI') AS 마지막_기록,
       COUNT(*) AS 건수
  FROM audit_log WHERE target_type = 'server' OR action LIKE 'SERVER\\_%' GROUP BY 1 ORDER BY 1;
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 월, action AS 종류, COUNT(*) AS 건수
  FROM audit_log WHERE action LIKE 'ADMIN\\_LOGIN%' OR action = 'ADMIN_LOGOUT' GROUP BY 1, 2 ORDER BY 1, 2;
"""), '10-21 실행분', '07/31 ~ 09/26 공백 있음(물어보면 답 시트) · 10-21 실측 = SERVER_LOGIN 마지막 10-21 10:04 · COMMAND · AUDITD 11:09 · HEARTBEAT 11:04 = 살아 있음')

add('G-7', '3.4 ⑥ · 3.5 ⑤', '발송 요청 인증수단(토큰 · 설치 묶음) · 고객사 허용 발신번호 관리', G_INVITO, pg(GW_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, a.target_type AS 대상, u.username AS 처리자, a.target_id AS 대상_식별값
  FROM audit_log a LEFT JOIN admin_user u ON u.id = a.admin_user_id
 WHERE a.action IN ('ROTATE_TOKEN','ISSUE_INSTALL_BUNDLE','SAVE_INSTALL_APPROVAL','AGENT_CREDENTIAL_ACK','UPSERT_CUSTOMER_SENDER_NUMBER',
                    'BULK_UPSERT_CUSTOMER_SENDER_NUMBER','UPDATE_CUSTOMER_SENDER_NUMBER','DEACTIVATE_CUSTOMER_SENDER_NUMBER','SET_SENDER_NUMBER_POLICY')
 ORDER BY a.created_at DESC LIMIT 40;
SELECT verification_status AS 승인_상태, is_active AS 사용, COUNT(*) AS 번호_수, MIN(created_at) AS 첫_등록, MAX(created_at) AS 최근_등록
  FROM customer_sender_number GROUP BY 1, 2 ORDER BY 1, 2;
"""), OLD)

add('G-8', '2.1 · 4.1 ②', '계정 최초 발급 기록(고객사 등록 · 발송 계정 생성 · 연결 마법사)', G_INVITO, pg(GW_PSQL, """
SELECT a.created_at AS 일시, a.action AS 종류, a.target_type AS 대상, u.username AS 발급자, a.target_id AS 대상_식별값
  FROM audit_log a LEFT JOIN admin_user u ON u.id = a.admin_user_id
 WHERE (a.action = 'CREATE' AND a.target_type IN ('reseller', 'sender_account')) OR a.action = 'CONNECT_WIZARD'
 ORDER BY a.created_at;
"""), OLD)

add('G-9', '3.2 · 4.1 ⑤ · 4.2 ⑤', '관리자 계정과 감사 기록 조회 권한', G_INVITO, pg(GW_PSQL, """
SELECT username AS 계정, role AS 등급, is_active AS 사용, last_login_at AS 마지막_로그인 FROM admin_user ORDER BY username;
"""), OLD, '감사 기록 조회 = admin 계정 하나')

add('G-10', '4.3', '월간 로그 점검 자료', G_INVITO, pg(GW_PSQL, """
SELECT to_char(date_trunc('month', created_at), 'YYYY-MM') AS 점검_대상_월,
       COUNT(*) FILTER (WHERE action = 'ADMIN_LOGIN_MFA') AS 관리자_로그인_완료, COUNT(*) FILTER (WHERE action = 'ADMIN_LOGIN_FAILED') AS 관리자_로그인_실패,
       COUNT(*) FILTER (WHERE action = 'ADMIN_LOGIN_BLOCKED') AS 관리자_로그인_차단, COUNT(*) FILTER (WHERE action = 'HARD_DELETE') AS 삭제,
       COUNT(*) FILTER (WHERE action = 'ROTATE_TOKEN') AS 토큰_재발급,
       COUNT(*) FILTER (WHERE action IN ('TOPUP', 'ADJUST_BALANCE', 'WALLET_TRANSFER')) AS 잔액_변경
  FROM audit_log WHERE created_at >= '2026-09-01' GROUP BY 1 ORDER BY 1;
SELECT to_char(date_trunc('month', first_seen_at), 'YYYY-MM') AS 점검_대상_월, SUM(event_count) FILTER (WHERE outcome = 'ok') AS 접속_성공,
       SUM(event_count) FILTER (WHERE outcome = 'denied') AS 접속_거부, SUM(event_count) FILTER (WHERE outcome = 'error') AS 접속_오류,
       COUNT(DISTINCT remote_ip) FILTER (WHERE outcome = 'denied') AS 거부된_출발지_수
  FROM account_access_event WHERE first_seen_at >= '2026-09-01' GROUP BY 1 ORDER BY 1;
"""), '10-21 실행분(고친 로그인 칸 확인)', '9월 완료 88 · 10월 49 · 10-21 이전 판은 ADMIN_LOGIN 으로 세서 0 이 나왔다(실제 기록 이름 ADMIN_LOGIN_MFA) = D-32')

add('G-11', '4.1', '발송 이력 — 보관 범위 · 최근 하루 집계 · 최근 원문 10건', G_INVITO, pg(GW_PSQL, """
SELECT MIN(created_at) AS 가장_오래된_발송_기록, MAX(created_at) AS 가장_최근_발송_기록 FROM message_request;
SELECT agent_id AS 발송_계정, msg_type AS 유형, norm_status AS 상태, COUNT(*) AS 건수
  FROM message_request WHERE created_at >= now() - interval '1 day' GROUP BY 1, 2, 3 ORDER BY 1, 2, 3;
SELECT id, created_at AS 요청, agent_id AS 발송_계정, left(phone, 3) || '****' || right(phone, 4) AS 수신번호, callback AS 발신번호,
       msg_type AS 유형, norm_status AS 상태, report_code_raw AS 결과_코드, report_time AS 결과_시각
  FROM message_request ORDER BY id DESC LIMIT 10;
"""), '10-21 실행분', '4.1 발송 이력은 이것으로 보여 준다 · 전체 집계는 큰 표를 다 훑어서 하지 않는다(날짜 범위로만) · 10-21 실측 = 2026-06-15 ~ 현재')

add('G-12a', '3.1 ①', '방화벽 정책 · 서비스(방화벽 · 침입 차단 · 감사 수집) 가동 · 정책 파일 수정 시각', G_ROOT,
    "ufw status numbered; systemctl is-active ufw fail2ban auditd; stat -c '%y  %n' /etc/ufw/user.rules /etc/ufw/after.rules", OLD,
    '9443 LIMIT Anywhere 줄 = 출처 판정은 앱 허용 목록(G-3)이 한다 · 둘을 한 쌍으로 보여 준다')
add('G-12b', '3.1 ④', '방화벽 변경 명령 기록', G_ROOT,
    "zgrep -h -a 'COMMAND=.*ufw ' /var/log/auth.log* 2>/dev/null | sort | tail -60", OLD,
    '09/24 변경은 이 기록에 없다 → 운영 기록 대장(H-16)으로 보완')
add('G-12c', '3.1 ④', '서버 관리자 접속 이력(최근 30건)', G_ROOT, 'last -n 30 -F -i | head -34', OLD)
add('G-12d', '3.1 ①', '외부로 열린 포트', G_ROOT, "ss -tlnH | awk '{print $4}' | sort -u", OLD)
add('G-12e', '3.1 ① · 4.1 ⑤', 'DB 컨테이너가 서버 안(127.0.0.1)에만 열림', G_INVITO,
    "docker ps --format 'table {{.Names}}\\t{{.Status}}\\t{{.Ports}}'", '10-21 실행분',
    '이름과 달리 bito-bench-postgres 가 운영 DB · 10-21 실측 = postgres · redis 둘 다 127.0.0.1')

add('G-12f', '3.1 ③', '사설 원격 접속망(Tailscale)이 꺼져 있는가', G_INVITO,
    'systemctl is-active tailscaled; systemctl is-enabled tailscaled', '10-21 실행분(점검 때 active · enabled → 10-21 두 서버 끔 · 대장 4건 = D-19 닫힘)',
    '꺼져 있으면 inactive · disabled · 10/03 해제 → 10/06 출장 때 다시 켬 → 귀국 뒤 끈 것까지 대장에 있어야 한다')

add('G-13a', '4.1 ⑤ · 4.2', '시스템 저널 · 웹 서버 로그 1년 보관 설정 · 가장 오래된 기록', G_ROOT,
    "grep -E '^(SystemMaxUse|MaxRetentionSec)' /etc/systemd/journald.conf; journalctl --disk-usage; journalctl --no-pager -o short-iso | head -1; "
    "grep -E '^\\s*(daily|weekly|rotate|compress)' /etc/logrotate.d/nginx", OLD, '설정일 2026-10-03 · 그 전은 08/22 이후분만 남아 있다')
add('G-13b', '2.2 ① · 4.2', '엔진 로그의 미등록 출발지 거부 줄(최근 30일 · 마지막 5줄)', G_ROOT,
    "journalctl -u bito-gateway --no-pager -o short-iso --since '30 days ago' | grep 'IP allowlist denied' | tail -5", OLD)
add('G-13c', '4.2', 'DB 서버 로그(최근 24시간 · 마지막 30줄)', G_INVITO,
    'docker logs --since 24h bito-bench-postgres 2>&1 | tail -30', '10-21 실행분',
    '10-21 실측 = 체크포인트 기록 + 10:02 lock timeout 1건(Q-87)')

add('G-14', '4.1 ⑤', 'DB 백업 보관분 · 마지막 성공 · 실행 기록 · 감시 기록 · 경보 설정 여부', G_ROOT,
    "ls -lt /opt/bito-gateway/backups | sed -n '2,10p'; stat -c '%y  LAST_SUCCESS' /opt/bito-gateway/backups/LAST_SUCCESS; "
    "tail -10 /opt/bito-gateway/backups/cron.log; tail -6 /opt/bito-gateway/backups/monitor.log; grep -c '^ALERT_CMD=.\\+' /opt/bito-gateway/backups/.env",
    OLD, '매일 03:30 암호화 · 외부 서버 전송 · 감시 08:30')

# ───────────────────────────── 핵심 검색어(Ctrl+F) ─────────────────────────────
# 심사위원이 쓸 말 + Harold 가 떠올릴 말. 명령집 모든 줄에 하나씩(빠지면 생성이 멈춘다).
KEYWORDS = {
    'H-1': '감사 기록 보관 · 보관 기간 · 1년 · 가장 오래된 기록 · 월별 건수 · 삭제 불가',
    'H-2': '로그 종류 · 기록 종류 · 첫 기록일 · 어떤 로그가 남나',
    'H-3': '로그인 기록 · 로그인 실패 · 로그인 차단 · 5회 실패 · 동시 접속 · 로그아웃 · 접속 로그',
    'H-4': '페이지 접속 · 화면 접속 · 접속 기록 · 메뉴 접속',
    'H-5': '개인정보 조회 · 개인정보 수정 · 다운로드 · 내보내기 · 엑셀 다운로드 · 고객 삭제 · 개인정보 처리',
    'H-6': '회원가입 · 계정 발급 · 발급 기록 · 고객사 등록 · 계정 생성일',
    'H-6b': '월별 계정 · 4월 · 7월 · 일괄 등록 · 계정 유형 · 계정 몰림',
    'H-6c': '시연 회사 · 삭제 이력 · 한줄 뷰티 · 지워진 회사',
    'H-6d': '시연 회사 · 삭제 백업 · 백업 파일',
    'H-7': '본인확인 · 한국모바일인증 · KMC · 담당자 인증 · 중복가입 · 실명',
    'H-7b': '수신 번호 일치 · 본인확인 번호 · 번호 바꿔치기 · 수동 변경',
    'H-8': '발신번호 등록 · 발신번호 심사 · 서류 · 반려 사유 · 발신번호 배정 · 회선 수 · 회선 정책',
    'H-8b': '시험 번호 · 02-1234-5678 · 중복 승인 · 이새 · 발신번호 중복',
    'H-8c': '발신번호 삭제 이력 · 발신번호 변경 이력 · 배정 해제 · 배정 변경 · 등록 · 수정 · 삭제 기록',
    'H-9': '국외 차단 · 해외 IP · 국내 대역 · 국외 접속 · VPN · 프록시',
    'H-9b': '10월 7일 차단 · 국내 이용자 차단 · 오차단 · 대역 갱신 · 대역 교체 이력 · 예외 통과 · 감지',
    'H-10': '허용 IP · 예외 승인 · 화이트리스트 · 승인자 · 허용 기간 · 만료',
    'H-11': 'API 출발지 · 싱크에이전트 · 연동 접속 · 모든 접근 경로 · 기계 접속',
    'H-11b': '미등록 출발지 · 싱크 미등록 · 연동 차단 전 확인 · 10월 연동',
    'H-12': '부정사용 · 입금 명의 · 보류 · 소명 · 결제 · 입금자 · 수동 충전 · 우회 충전',
    'H-13': '다중인증 · MFA · 2차 인증 · 인증번호 · 잠금 · 인증 실패 · 번호 변경',
    'H-13b': '번호 변경 상세 · 수동 변경 · 본인확인 뒤 변경 · gwchae',
    'H-14': '추가인증 · 발신 인증 · 발송 인증 · 발신번호 인증',
    'H-15': '직원 계정 · 관리자 계정 · 권한 · 등급 · 권한 변경 · 퇴사 · 비밀번호 변경 · OTP 등록',
    'H-16': '운영 기록 대장 · 로그 점검 기록 · 방화벽 변경 대장 · 접근권한 점검 · 확인자',
    'H-17': '월간 점검 · 로그 점검 · 월 1회 · 점검 자료 · 로그인 실패 많은 계정',
    'H-18': '이용 정지 · 계정 제한 · 해지 · 스팸 신고 조치 · 고지 · 조치 결과',
    'H-18b': '고객사 해지 · 해지 사유 · 해지 고지 · 안내 문자 · 계약 종료 · 해지 기록',
    'H-19': '금칙어 · 스팸 키워드 · 차단 규칙 · 탐지 · 문안 차단',
    'H-21a': '방화벽 · ufw · 방화벽 정책 · 방화벽 상태',
    'H-21b': '방화벽 변경 이력 · 정책 변경 · 변경 명령',
    'H-21c': '서버 접속 이력 · 관리자 접속 · SSH 접속 · last',
    'H-21d': '열린 포트 · 포트 · 외부 노출',
    'H-21e': 'DB 외부 차단 · 도커 · docker · 127.0.0.1 · DB 포트',
    'H-21f': '침입 차단 · fail2ban · SSH 무차별 · 반복 실패 차단',
    'H-21h': '결제 통계 수신 DB · 23388 · DOCKER-USER · 컨테이너 방화벽 · 외부 열린 DB',
    'H-21g': '사설 원격 접속망 · Tailscale · 우회 경로 · 비가시 경로 · 출장',
    'H-22a': '앱 로그 · 애플리케이션 로그 · 로그 회전 · pm2 · 로그 권한',
    'H-22b': '웹 서버 로그 · nginx · 시스템 로그 · 저널 · 1년 보관 설정',
    'H-22c': 'DB 로그 · 데이터베이스 로그 · 장애 로그 · 오류 로그',
    'H-22d': 'DB 로그 설정 · 컨테이너 로그',
    'H-23': '백업 · 백업 주기 · 백업 실패 경보 · 암호화 백업 · 외부 보관',
    'H-24': '시행 스위치 · 적용 설정 · 적용 범위 · 시행일 · 10월 26일',
    'G-1': '게이트웨이 감사 기록 · 엔진 로그 보관 · 보관 기간',
    'G-2': '게이트웨이 로그 종류 · 엔진 기록 종류',
    'G-3': '게이트웨이 허용 IP · 에이전트 허용 IP · 지정 IP · 허용 사유 · 만료일',
    'G-4': '게이트웨이 접속 이력 · 에이전트 접속 · 접속 거부 · 미등록 출발지 · API 접속',
    'G-4b': 'API 미등록 출발지 · API 차단 전 확인 · 9월 거부 · 거부 계정',
    'G-5': '허용 IP 변경 이력 · 에이전트 변경 · 처리자',
    'G-6': '게이트웨이 서버 접속 기록 · 명령 기록 · 관리 사이트 로그인 · 관리자 로그인',
    'G-7': '토큰 재발급 · 인증키 · 인증서 · 설치 묶음 · 엔진 발신번호 · Agent 인증',
    'G-8': '게이트웨이 계정 발급 · 연결 마법사 · 리셀러 등록 · 발송 계정 생성',
    'G-9': '게이트웨이 관리자 계정 · 감사 기록 조회 권한 · 관리자 등급',
    'G-10': '게이트웨이 월간 점검 · 엔진 점검 자료',
    'G-11': '발송 이력 · 발송 로그 · 발송 결과 · 수신번호 · 결과 코드',
    'G-12a': '게이트웨이 방화벽 · 엔진 방화벽 · auditd · 9443',
    'G-12b': '게이트웨이 방화벽 변경 이력',
    'G-12c': '게이트웨이 서버 접속 이력 · 게이트웨이 SSH',
    'G-12d': '게이트웨이 열린 포트',
    'G-12e': '게이트웨이 DB 외부 차단 · 게이트웨이 도커 · docker',
    'G-12f': '게이트웨이 사설 원격 접속망 · 게이트웨이 Tailscale · 우회 경로',
    'G-13a': '게이트웨이 저널 · 게이트웨이 nginx · 게이트웨이 서버 로그 1년',
    'G-13b': '엔진 로그 · 미등록 출발지 거부 · IP allowlist denied',
    'G-13c': '게이트웨이 DB 로그 · 엔진 DB 로그',
    'G-14': '게이트웨이 백업 · 엔진 백업 · 백업 경보',
}
_ids = [r[0] for r in ROWS]
assert set(KEYWORDS) == set(_ids), f'검색어 빠짐/남음: {set(_ids) ^ set(KEYWORDS)}'
CMD_ROWS = [(no, item, KEYWORDS[no], title, where, cmd, basis, memo) for no, item, title, where, cmd, basis, memo in ROWS]

# ───────────────────────────── 다른 시트 ─────────────────────────────
REHEARSAL = [
    ('R-1', H_ADMIN, '새 조회가 쓰는 한줄로 표 컬럼 확인', '(실행 완료 10-21)', '10개 표 · 쓰는 컬럼 전부 있음(Harold 실행 출력)'),
    ('R-2', G_INVITO, '새 조회가 쓰는 게이트웨이 표 컬럼 확인', '(실행 완료 10-21)', 'agent_account · message_request · 쓰는 컬럼 전부 있음(Harold 실행 출력)'),
    ('R-3', '-', '「명령집」 시트를 위에서부터 한 번씩', '-', '오류 난 번호만 알려 주면 그 줄만 고친다'),
]

FIND = [
    ('2.1 계정관리(발급 · 본인확인 · 발신번호 · 회선)', 'H-6 · H-6b · H-6c · H-6d · H-7 · H-8 · H-8c(변경 · 삭제 이력)', 'G-8'),
    ('2.2 국외 이용자(차단 · 예외 · 모든 경로)', 'H-9 · H-10 · H-11 · H-24', 'G-3 · G-4 · G-5 · G-13b'),
    ('2.3 부정사용(입금 명의 보류 · 소명)', 'H-12', '-'),
    ('3.1 방화벽(정책 · 변경 이력 · 관리자 접속)', 'H-16 · H-21', 'G-6 · G-12'),
    ('3.2 · 3.3 접근권한(등급 · 변경 · 점검)', 'H-15 · H-16', 'G-9'),
    ('3.4 다중인증', 'H-7 · H-13 · H-24', 'G-7'),
    ('3.5 추가인증(발신 인증)', 'H-14 · H-24', 'G-7'),
    ('4.1 이용자 로그(발송 이력 = G-11)', 'H-1 ~ H-7', 'G-1 · G-2 · G-4 · G-11'),
    ('4.2 시스템 로그(관리 · 개인정보 · 서버 · DB)', 'H-5 · H-15 · H-22', 'G-1 · G-13'),
    ('4.1 ⑤ · 4.2 ⑤ 삭제 · 변조 방지 · 백업', 'H-1 · H-21e · H-22 · H-23', 'G-9 · G-12e · G-13 · G-14'),
    ('4.3 로그 점검(월 1회 · 기록 · 후속조치)', 'H-16 · H-17', 'G-10'),
    ('5.1 조치 및 결과 확인', 'H-18 · H-18b(해지 사유 · 안내 문자)', '-'),
    ('5.2 금칙어 · 악성 URL(KISA)', 'H-19(금칙어 규칙)', '링크가드 관제센터 → URL 정책(출처 KISA) · 판정 이력 · 운영 현황(KISA 마지막 적재) = 대표 화면'),
    ('화면으로 보여 주기(가림이 이미 들어가 있다)', '슈퍼관리자 sys.hanjullo.com → 보안 · 인증: 감사 로그 · 국외 접근 이력 · 운영 기록 대장 · 본인인증 현황 · 로그인 차단 · 금칙어 차단 · 직원 계정 · 권한', '게이트웨이 관리 콘솔 → 감사 기록 · Agent 수정 창(허용 IP 사유 · 만료일)'),
]

# (번호, 심사 항목, 핵심 검색어, 예상 질문, 답변 = 그대로 말할 문장, 보여 줄 번호, 근거 · 주의)
# 모의 실사(2026-10-21) 질의마다 여기에 쌓는다. ⚠ 표시 = 아직 확인 안 된 것 · Harold 가 답할 것.
QA = [
    ('Q-01', '2.1', '회원가입 · 계정 발급 · 발급 기록',
     '계정 발급(회원가입) 기록을 보여 주십시오.',
     '「개인 · 외국인 가입은 받지 않고, 법인 계약 뒤 당사가 계정을 발급합니다. 발급할 때마다 발급자 · 일시 · 대상 계정 · 고객사가 감사 기록에 남습니다.」',
     'H-6', '모의 질의 1 · 10-21 실측 15줄'),
    ('Q-02', '2.1', '10월 2일 이전 계정 · 생성일',
     '10월 2일 전에 만든 계정은 무엇으로 확인합니까?',
     '「발급 기록 장치는 10월 2일에 넣었습니다. 그 전 계정은 계정 생성일로 확인합니다.」',
     'H-6 둘째 표 · H-6b', '10/02 이전 발급 기록이 있다고 말하지 않는다'),
    ('Q-03', '2.1', '발급자 빈칸 · 스크립트 · 시연 회사',
     '발급자 칸이 빈 줄은 누가 발급한 것입니까?',
     '「당사 시연 · 시험용 회사(한줄 뷰티(시연) · 미가입 테스트)를 서버 관리 스크립트로 만든 것이라 발급자 칸이 비었습니다. 고객사 발급은 전부 발급자가 남습니다. 10월 21일부터는 스크립트로 만든 계정 · 회사도 경로(서버 관리 스크립트) · 스크립트 이름 · 실행한 서버 계정이 함께 남습니다.」',
     'H-6', '근거 = 10-21 배포 S2(account-issue.ts channel script · osUser) · 지난 기록은 바꾸지 않는다 · ⚠ 지난 스크립트를 누가 돌렸는지는 Harold 가 답'),
    ('Q-04', '2.1', '시스템 계정 · system_sync · 연동 계정',
     'system_sync_ 로 시작하는 계정은 무엇입니까? 사람이 쓸 수 있습니까?',
     '「고객 데이터 연동(싱크) 기록용 시스템 계정입니다. 로그인이 막혀 있고, 로그인을 시도하면 차단 기록이 남습니다.」',
     'H-6 · H-3', '근거 = routes/auth.ts:297(시스템 계정 로그인 403 + 로그인 차단 기록)'),
    ('Q-05', '2.1', '4월 · 7월 계정 몰림 · 일괄 등록',
     '4월 153개 · 7월 53개는 왜 몰렸습니까?',
     '「4월은 기존 시스템에서 한줄로로 옮기면서 고객사마다 관리자 계정을 일괄 생성했고, 7월은 발송 에이전트를 쓰는 고객사가 발송 통계를 볼 수 있도록 관리자 계정을 일괄 생성했습니다. 4월은 관리자 66개가 고객사 66곳, 7월은 53개가 53곳으로 회사마다 하나씩입니다.」',
     'H-6b', '10-21 실측 · 이유 = Harold 10-21(4월 레거시 이전 일괄 생성 · 7월 에이전트 통계 확인용 일괄 생성)'),
    ('Q-06', '2.1', '시연 회사 삭제 · 삭제 이력 · 두 번 등록',
     '「한줄 뷰티(시연)」가 같은 날 두 번 등록됐고 첫 회사는 지금 없습니다. 삭제 이력을 보여 주십시오.',
     '「당사 내부 시연용 회사입니다(고객사 아님). 10월 9일 새로 만들면서 옛 회사를 은퇴시켰고, 10월 10일 대표 지시로 서버 정리 스크립트로 완전 삭제했습니다. 고객사는 이 경로로 지울 수 없습니다(시연 표식이 있는 회사만 대상). 그때 정리는 감사 기록에 남지 않아, 10월 21일 서버 관리 스크립트도 실행자와 삭제를 감사 기록에 남기도록 보완했습니다.」',
     'H-6c', '근거 = purge-retired-demo-company.ts(대상 조건) · 10-21 실측 = 감사 기록에 삭제 줄 없음 · 삭제 백업 파일 administrator · root 홈 둘 다 없음(지난 판 「백업 파일로 남겼다」 문장 삭제) · 마지막 문장 = 10-21 배포 S2(삭제와 같은 트랜잭션에 company_purged 기록)'),
    ('Q-07', '2.1 ②', '본인확인 · 한국모바일인증 · 담당자 인증',
     '본인확인 이력을 보여 주십시오.',
     '「계정 담당자는 첫 로그인 때 한국모바일인증으로 본인확인을 합니다. 10월 6일 운영을 시작해 직원 5계정이 마쳤고, 전 계정은 10월 26일부터입니다. 인증 이름 · 번호 · 요청번호 · 접속 IP가 남습니다.」',
     'H-7', '모의 질의 2 · 10-21 실측 11줄'),
    ('Q-08', '2.1 ②', '본인확인 대기 · pending',
     '상태가 pending 인 건(10/14 psy5868)은 무엇입니까?',
     '「담당자 변경을 시작하고 인증을 끝내지 않은 건이라 번호는 바뀌지 않았습니다.」',
     'H-7', '근거 = 번호는 verified 로 확정될 때만 바뀐다(SCHEMA.md identity_verifications)'),
    ('Q-09', '2.1 ②', '본인확인 무효 · superseded',
     'superseded 는 실패입니까?',
     '「아닙니다. 같은 계정이 인증을 다시 시작하면 앞의 대기 건을 무효로 돌린 것입니다.」',
     'H-7', '같은 근거'),
    ('Q-10', '2.1 ②', '고객사 본인확인 · 10월 26일 · 전 계정 적용',
     '고객사 계정의 본인확인 기록이 없는데요?',
     '「직원 계정부터 운영을 시작했고, 전 계정 적용은 2026년 10월 26일입니다. 로그인 화면에서 미리 고지하고 있습니다.」',
     'H-7 · H-24 · 로그인 화면 안내 창', '문서 2.1 3절 「2026-10-06 운영 개시 · 2026-10-26 전 계정」과 같은 말 · 10-21 실측 H-24 = 본인확인 적용 명단 직원 5계정(hoyun · gwchae · sgbaek · psy5868 · mobile) · 10/26 = 그 명단을 전 계정(*)으로 바꾼다(설정 한 줄 · 자동 전환 아님) · 「그날 무엇이 바뀌나」에는 이 문장'),
    ('Q-11', '2.1 ②', '중복가입 · 중복 계정 · 중복확인값',
     '중복가입은 어떻게 확인합니까?',
     '「본인확인기관이 주는 중복가입 확인값을 원문 없이 해시로 저장합니다. 같은 사람이 계정을 더 만들면 같은 값으로 찾습니다. 계정 발급은 당사 한 창구뿐이라 두 번째 계정도 당사 발급 절차를 거칩니다.」',
     'H-7(중복확인값_기록 칸)', '10-21 실측 = 인증 완료 6건 전부 t · 대기 · 무효 건은 f(인증 전이라 값이 없다) · ⚠ 「차단한다」고 말하지 않는다(문서 = 기록 + 발급 창구 통제)'),
    ('Q-12', '4.1 · 4.2', '감사 기록 보관 · 삭제 불가 · 한줄로',
     '한줄로 감사 기록은 언제부터 있습니까? 지울 수 있습니까?',
     '「2026년 2월 11일부터 남아 있고, 감사 기록을 지우는 기능은 시스템에 없습니다. 관리자 화면도 조회만 됩니다.」',
     'H-1', '10/03 수집 기준(내일 건수는 그 뒤로 늘어 있다)'),
    ('Q-13', '4.1 · 4.2', '감사 기록 보관 · 게이트웨이 · 엔진',
     '게이트웨이 감사 기록은 언제부터입니까?',
     '「2026년 6월 15일부터 남아 있고, 지우는 코드가 없습니다.」',
     'G-1', '10/03 수집 기준'),
    ('Q-14', '3.1 ④', '서버 접속 기록 공백 · 명령 수집 · 7월 31일',
     '게이트웨이 서버 접속 · 명령 기록에 비는 기간이 있습니다.',
     '「7월 31일부터 9월 26일까지 수집이 멈춰 있었습니다. 9월 점검에서 찾아 재개했고 그 뒤로 연속입니다. 그 기간에 기록이 있다고 하지 않습니다.」',
     'G-6 · H-16', '공백 = 07/31 16:33 ~ 09/26 22:31(10/03 수집)'),
    ('Q-15', '4.2', '저널 8월 22일 · 서버 로그 1년',
     '게이트웨이 서버 저널이 8월 22일부터입니다.',
     '「1년 보관 설정을 10월 3일에 넣었습니다. 그 전에는 기본 용량 한도로 잘려 8월 22일 이전분은 없습니다.」',
     'G-13a', '10/03 수집'),
    ('Q-16', '4.3 ③', '접속 거부 급증 · 이상징후 · 후속 조치',
     '10월 초 게이트웨이 접속 거부가 많습니다.',
     '「사내 시험 서버의 옛 시험 발송 에이전트가 허용 IP를 바꾼 뒤에도 옛 출발지로 재접속한 것입니다. 외부 공격이 아닙니다. 10월 3일 12시 12분에 정지했고 다시 조회해 멈춘 것을 확인했습니다. 이상징후 후속 조치 사례입니다.」',
     'G-4', '10-21 실측 = 10월 에이전트 거부 2,252건 · 마지막 10/03 12:12:40(bito-test-97 · 58.227.193.66) · 그 뒤 거부 0'),
    ('Q-17', '2.2', '국외 차단 · 시작일 · 프록시 예외',
     '국외 차단은 언제부터입니까?',
     '「탐지는 8월 28일, 차단은 10월 1일부터입니다. 막힌 것 중 고객사 보안 프록시 출구는 사유와 승인자를 붙여 예외로 등록했습니다.」',
     'H-9 · H-10', '10/03 수집'),
    ('Q-18', '4.3', '9월 로그 점검 · 점검 기록 · 운영 기록 대장',
     '9월 로그 점검은 언제 했습니까?',
     '「점검 기록 대장이 생긴 10월 3일 뒤 처음 했습니다. 10월 4일 지원팀장이 작성하고 10월 5일 대표가 확인했습니다. 날짜를 앞당겨 적지 않았습니다.」',
     'H-16', '대장 조회 10/05'),
    ('Q-19', '3.4', '구글 OTP · 관리자 OTP · 다중인증 구분',
     '구글 OTP는 무엇입니까? 고객 다중인증입니까?',
     '「구글 OTP는 당사 직원이 관리자 사이트에 들어갈 때 쓰는 것입니다. 고객 다중인증은 휴대폰 인증번호입니다.」',
     'H-13 · H-15', ''),
    ('Q-20', '4.1', '발송 이력 · 발송 로그',
     '발송 이력은 어디서 봅니까?',
     '「비토 엔진 발송 이력으로 보여 드리겠습니다.」',
     'G-11', '⚠ 외주 중계 라인(발송 DB 1~11)은 보여 주지 않는다(★10-21 Harold)'),
    ('Q-21', '2.1 ⑤', '발신번호 신청 · 심사 · 서류 · 반려 사유',
     '발신번호 등록 신청과 심사 이력을 보여 주십시오.',
     '「고객사가 증빙 서류를 붙여 신청하면 당사 직원이 서류를 확인해 승인하거나 사유를 적어 반려합니다. 신청자 · 신청 일시 · 서류 · 심사자 · 심사 일시 · 반려 사유가 건마다 남습니다. 서류가 맞지 않으면(다른 번호의 통신이용증명원 · 위임장 누락) 반려하고 다시 받습니다.」',
     'H-8', '모의 질의 3 · 10-21 실측 26건'),
    ('Q-22', '2.1 ⑤', '반려 사유 이상한 글자 · 테스트계정 · 시험 신청',
     '테스트계정 반려 사유가 의미 없는 글자입니다.',
     '「당사 시험 고객사에서 회선 상한 거부를 시험한 신청이라 반려했고, 사유 칸에 시험 입력이 남았습니다. 기록은 고치지 않고 그대로 둡니다.」',
     'H-8', '근거 = 10/05 회선 상한 시험(작업 기록)'),
    ('Q-23', '2.1 ⑤', '시험 번호 · 02-1234-5678 · 인비토',
     '02-1234-5678 같은 번호가 증빙을 거쳐 승인됐습니까?',
     '「당사 회사(인비토)에서 8월 20일 신청 · 심사 기능을 시험한 번호입니다. 지금은 발신번호에서 빠져 있어 발송에 쓸 수 없습니다.」',
     'H-8 · H-8b', '10-21 실측 = 발신번호 0건 · ⚠ 「언제 지웠나」는 답할 기록이 없다(삭제 기록 없음 = Q-26)'),
    ('Q-24', '2.1 ⑤', '같은 번호 두 번 승인 · 이새 · 052-937-6647',
     '같은 번호가 8월 18일 · 19일 두 번 승인됐습니다.',
     '「같은 번호가 두 개로 등록되지는 않았습니다. 지금 등록은 8월 19일 승인분 하나입니다.」',
     'H-8 · H-8b', '10-21 실측 = 1건(등록 시각 = 8/19 승인 시각과 같음) · ⚠ 8/18 승인분이 왜 없는지(다시 신청한 이유)는 미확인 · 삭제 기록 없음 = Q-26'),
    ('Q-25', '2.1 ⑤', '발신번호 배정 · 계정별 번호 · 지정 사용',
     '번호를 계정별로 나눠 쓰는 이력을 보여 주십시오.',
     '「번호마다 사용 범위를 전체 또는 지정으로 정하고, 지정이면 고른 계정만 그 번호로 보냅니다. 고객사 관리자가 자기 회사 계정에 배정하고, 배정한 사람과 일시가 남습니다.」',
     'H-8 둘째 표', '10-21 실측 30건'),
    ('Q-26', '2.1', '발신번호 삭제 이력 · 배정 해제 · 삭제 기록',
     '발신번호를 지우거나 배정을 풀면 이력이 남습니까?',
     '「10월 21일부터 발신번호 등록 · 수정 · 기본 번호 지정 · 삭제와 배정 범위 변경 · 배정 교체 · 배정 해제가 모두 감사 기록에 남습니다(처리자 · 고객사 · 번호 · 변경 전후 · 들어온 경로). 그 전 삭제 · 해제는 기록이 없었고, 점검에서 찾아 보완했습니다.」',
     'H-8c', '근거 = 10-21 배포 S1(utils/change-audit.ts) · 실측 10-21 13:15 등록 → 삭제 두 줄(ceo · 테스트계정 · 02-1234-5678) = H-8c 로 보여 준다'),
    ('Q-27', '2.1 ③ · ④', '회선 정책 · 법인 일괄 지정 · SQL 처리 · 전후 값',
     '회선 정책을 DB에서 직접 바꿨습니까? 바뀐 게 없는 기록도 있습니다.',
     '「전 고객사를 법인으로 일괄 지정할 때 운영 DB에서 한 번에 처리했고, 같은 사유 문장을 감사 기록으로 남겼습니다. 값을 바꾸지 않고 저장해도 기록이 남습니다(전후 값이 같은 줄).」',
     'H-8 셋째 표', '10-21 실측 · 테스트계정 유선 상한 = 10/05 21:00 시험값 1 → 10/21 13:15 대표가 화면에서 5(회선 정책 변경 기록 2줄 · 시험 회사)'),
    ('Q-28', '3.2', 'admin 계정 · 공용 계정',
     '심사자 admin 은 누구 계정입니까? 공용 계정입니까?',
     '「admin 은 여러 담당자가 함께 쓰던 관리자 계정이었습니다. 누가 한 일인지 특정할 수 없어 담당자마다 개인 계정을 발급하고, admin 은 사유를 남겨 사용 중지했습니다(문서 3.2 1-6절). 8월 20일 심사 기록은 그 정리 전의 기록입니다.」',
     'H-8 · H-15', 'Harold 10-21 = 공용 계정이라 사용 중지 · 하나씩 개인 계정 발급 · 10-21 실측 = admin 사용 중지 2026-08-27 23:18:28(ceo · 사유 「개인계정으로 전환」) · 같은 날 22:54 suran 개인 계정 발급(사유 「공용계정 admin 개인계정 전환」) → 8/20 심사는 중지 전 = 확정'),
    ('Q-29', '2.2 ①', '국내 대역 · 허용 대역 · APNIC · 2,597',
     '국내 허용 대역이 실제로 적재돼 있습니까?',
     '「아시아태평양 인터넷 주소 관리기관(APNIC)이 한국에 할당한 대역 2,597개를 넣어 두고, 여기에 없는 주소는 국외로 판정합니다.」',
     'H-9 첫 표', '모의 질의 4 · 10-21 실측 = 2,597개 · 마지막 갱신 10/07 16:07'),
    ('Q-30', '2.2 ④ · 6절', '프록시 차단 · 10월 1일 16건 · 룰루레몬',
     '10월 1일 차단 16건과 그 뒤 처리를 보여 주십시오.',
     '「고객사가 사내 보안 프록시를 거쳐 접속하면서 해외 출구로 나와 막혔습니다. 그 계정에만, 사유와 승인자를 적어 예외로 등록했고, 그 뒤 접속은 예외 통과로 기록됩니다.」',
     'H-9 셋째 표 · H-10', '10-21 실측 16건(167.103.97.x) · 예외 등록 = H-10'),
    ('Q-31', '2.2 ③ · 7절', '대표 출장 · 일본 · 예외 만료 · 허용 기간',
     '예외에 허용 기간을 둔 실제 사례가 있습니까?',
     '「10월 8일 대표 해외 출장 중 접속이 막혀, 사유 · 승인자 · 만료일(10월 12일)을 붙여 예외로 등록했습니다. 만료일이 지나면 통과가 끊깁니다.」',
     'H-9 셋째 표 · H-10', '10-21 실측 = 10/08 15:20 ceo 219.164.64.138 차단 → 16:23 예외(사유 · 승인자 ceo · 만료 10/12 23:59) · 대장 「유효 t」는 회수 여부 칸이고 만료는 판정 때마다 본다(utils/geo-access.ts:246 expires_at > NOW()) → Q-36 · 만료 뒤 그 주소 접속은 없었다(실제로 끊긴 장면은 없다 · 있다고 말하지 않는다)'),
    ('Q-32', '2.2 ①', '10월 7일 차단 · 국내 이용자 차단 · 오차단',
     '10월 7일 국내 주소로 보이는 이용자들이 국외로 막혔습니다. 판정이 틀린 것 아닙니까?',
     '「직원이 예외 승인을 하려다 국내 대역 일괄 교체 칸에 주소 하나를 넣어 2,597개가 1개로 바뀌었고, 15시 57분부터 16시 06분까지 국내 이용자 4계정이 국외로 막혔습니다. 같은 날 16시 07분에 원천 목록으로 되돌렸고, 그 교체 기능과 화면 칸을 없애 직원은 예외 승인만 쓰게 했습니다. 지금 그 주소들은 모두 국내 대역에 들어 있습니다. 경위는 운영 기록 대장에 남겼습니다.」',
     'H-9 · H-9b', '10-21 실측 = 교체 감사 기록 suran 10/07 15:47:51 「2597 → 1」 · 차단 15:57~16:06(espayment1·3·4 14.52.61.97 · mdysresort 61.81.163.98 · hoyun 115.138.27.202) · 대역 마지막 갱신 16:07:41 · 세 주소 지금 국내 = t · 경위 원문 = status/OPS.md 2-2-F ★2026-10-07(일괄 등록 카드 · POST /geo/cidrs/bulk 삭제 · 같은 날 SQL 로 복구 · 직원은 예외 승인만) · 10/03 빌드에는 경로가 있고 지금 소스에는 없다 · ⚠ 16:07 SQL 복구는 감사 기록이 없다 → 대장 기록으로 보완(Q-34) · ⚠ 마지막 문장은 대장 기록 뒤에만'),
    ('Q-33', '2.2 ①', '감지 · 예외 통과 · 차단 시행 뒤 통과',
     '차단을 시행한 뒤에도 매일 「감지」가 있습니다. 막지 않고 통과시킨 것입니까?',
     '「차단을 시행한 10월 1일 뒤의 감지는 전부 예외 승인을 받은 계정의 통과 기록입니다. 예외가 없는 국외 접속은 차단으로 기록됩니다.」',
     'H-9 둘째 표 · H-9b 셋째 표', '10-21 실측 = 10/02~10/19 감지 44건 전부 exempted=true'),
    ('Q-34', '4.3 ③ · 2.2', '10월 7일 · 이상징후 · 후속 조치 · 운영 기록 대장',
     '10월 7일 국내 이용자 차단을 점검에서 찾았고, 후속 조치 기록이 있습니까?',
     '「그날 바로 국내 대역을 되돌리고 교체 기능을 없앴습니다. 10월 21일 점검에서 이 건을 다시 확인해 운영 기록 대장(로그 점검)에 이상 · 원인 · 조치 · 재확인을 기록했습니다.」',
     'H-16', '대장 로그 점검 10-21 등록(서수란 · 대표 확인 · 작성일 10-21 · 소급 아님) = H-16'),
    ('Q-35', '2.2 ④', 'VPN 예외 · 시세이도 · shadmin',
     '예외 사유에 「VPN」이 있습니다. 우회 접속을 허용한 것입니까?',
     '「아닙니다. 고객사 본사가 업무상 반드시 거치게 하는 사내망의 고정 출구 주소 하나만, 그 계정 하나에만, 사유와 승인자를 적어 허용했습니다. 일반 VPN · 프록시는 국외 출구면 그대로 막힙니다.」',
     'H-10', '10-21 실측 = shadmin 153.72.39.239/32 · 10/07 15:22 차단 → 16:00 예외(ceo) · 그 주소는 지금도 국외 판정(H-9b f) = 예외로만 통과'),
    ('Q-36', '2.2 ③', '만료 지난 예외 · 유효 표시 · 자동 만료',
     '대표 출장 예외는 10월 12일에 만료됐는데 대장에 아직 「유효」로 나옵니다.',
     '「유효 칸은 회수했는지를 나타냅니다. 만료일은 접속을 판정할 때마다 따로 보기 때문에 10월 12일이 지나면 그 주소는 통과하지 않습니다.」',
     'H-10', '근거 = utils/geo-access.ts:246(expires_at IS NULL OR expires_at > NOW())'),
    ('Q-37', '2.2 ③', '전역 예외 · global · 대표 출장',
     '대표 출장 예외는 왜 계정이 아니라 전역 범위입니까?',
     '「대표가 쓰는 운영 관리자 계정은 고객 계정과 다른 표에 있어 계정 범위로 걸 수 없습니다. 그래서 그 출장지 주소 하나만, 4일 기한으로 허용했습니다.」',
     'H-10', '근거 = access_origin_allowlist.user_id 외래키가 users(SCHEMA 73번) · 슈퍼관리자 계정 범위는 추가 과제'),
    ('Q-38', '2.2 ③', '이새 · 같은 주소 두 줄 · 범위 착오',
     '이새에프앤씨는 같은 주소가 두 번 등록돼 있습니다.',
     '「8월 30일 등록이 회사 API 범위로 잘못 들어가 10월 3일 에이전트 범위로 다시 등록했습니다(사유에 적혀 있습니다). 쓰지 않는 회사 API 범위 옛 줄은 10월 21일 회수했고, 회수 사유는 운영 기록 대장에 남겼습니다.」',
     'H-10', '옛 회사 API 줄 10-21 13:49 회수(대표) · 회수 사유 = 대장 · H-10 둘째 표에 회수 기록 · D-15 · D-35'),
    ('Q-39', '2.2 ③', '예외 SQL 등록 · channel sql · 일괄 등록',
     '예외를 DB에서 직접 넣었습니까?',
     '「10월 2일 · 3일 여러 계정을 한 번에 등록할 때 운영 DB에서 처리했고, 같은 사유 · 승인자로 감사 기록을 함께 남겼습니다(channel: sql).」',
     'H-10 둘째 표', '10-21 실측'),
    ('Q-40', '2.2 ③', '만료 없는 예외 · 무기한 · 유효성 점검',
     '예외 대부분이 만료일이 없습니다.',
     '「고객사 사내 보안망 · 고정 서버처럼 상시 쓰는 출구라 만료를 두지 않았고, 월간 로그 점검에서 예외가 아직 필요한지 확인합니다(점검 항목 3번).」',
     'H-10 · H-16', '근거 = 4.1-4.3 문서 6-2 표 3번 「예외 승인의 유효성 · 만료」 · ⚠ 9월 점검 기록에 이 확인이 들어 있는지는 질의 4.3 때 확인'),
    ('Q-41', '2.2 ⑤', '연동 경로 · 기록만 · 모든 접근 경로 · 미등록 출발지',
     'SDK · 싱크에이전트 · API 연동은 미등록 출발지를 기록만 합니다. 이걸로 통제가 됩니까?',
     '「처음에는 미등록 출발지를 기록만 하면서 고객사마다 쓰는 출발지를 모두 등록받았고, 기록에서 미등록 출발지가 더 나오지 않는 것을 확인한 뒤 10월 21일부터 차단으로 바꿨습니다. 지금은 한줄로 SDK · 싱크에이전트와 발송 엔진 API 모두 등록된 출발지가 아니면 막힙니다.」',
     'H-11 · H-11b · G-4 · G-4b', '켠 시각 = 한줄로 10-21 14:24:18 재기동 · 게이트웨이 10-21 14:26 재기동(enforce) · 켠 뒤 게이트웨이 API 28건 ok · 차단 0 · 한줄로 감지 · 차단 0 · 10-21 실측 = 한줄로 회사 API 미등록 0· 싱크에이전트 10월 미등록 60건(2곳) · 게이트웨이 API 9,439건 출발지 1곳 · 사유별은 G-4b'),
    ('Q-42', '2.2 ① · 4.3', '9월 거부 · 에이전트 거부 15건',
     '9월 게이트웨이 접속 거부 15건(출발지 4곳)은 무엇입니까?',
     '「허용되지 않은 접속을 막은 기록입니다. 당사 규격연동 시험 계정 5건(9월 25일 · 인증서 없는 접속 2건 포함), 당사 한줄로 서버 발송 에이전트 1건(9월 24일 설치 다음 날), 고객사 에이전트 2곳 9건입니다.」',
     'G-4 · G-4b', '10-21 실측 = itensms03 1.201.116.44 8건(9/9 15:36~15:42) · SPECTEST01 115.138.27.202 5건(9/25) · hanjul-04 58.227.193.62 1건 · eternal01 3.39.204.6 1건(둘 다 9/24 07:03:06 같은 초) · ⚠ 고객사 2곳의 실패 원인은 확인 안 함(물으면 「인증 실패로 거부」까지만)'),
    ('Q-43', '2.2 ⑤', '규격연동 · 인증서 · SPEC_CLIENT_CERT_REQUIRED',
     '거부 사유 SPEC_CLIENT_CERT_REQUIRED 는 무엇입니까?',
     '「규격연동 접속은 클라이언트 인증서가 있어야 하고, 인증서 없이 들어온 접속 2건(9월 25일)을 거부한 기록입니다.」',
     'G-4', '10-21 실측 = 2건 · 출발지 1곳 · 9/25 06:55~07:25'),
    ('Q-44', '2.3', '입금 명의 보류 · 소명 · 처리 기록',
     '입금자 명의가 다른 건의 보류 · 소명 · 처리 기록을 보여 주십시오.',
     '「무통장 입금은 입금자명을 계정의 회사명 · 대표자명과 시스템이 대조하고, 다르면 자동으로 거절하지 않고 보류합니다. 보류 사유 · 시각 · 고객사 소명 · 처리자 · 처리 일시 · 결과가 신청 건과 감사 기록 두 곳에 남습니다.」',
     'H-12', '모의 질의 6 · 10-21 실측 보류 4건 · 처리 기록 4건'),
    ('Q-45', '2.3', '소명 없이 승인 · 47초 · 지삼유통',
     '10월 20일 30만 원은 소명도 없이 47초 만에 승인됐습니다. 무엇을 근거로 명의를 확인했습니까?',
     '「입금자명이 「(대표자 이름) 지삼유통」으로 등록된 회사명이 들어 있었습니다. 시스템은 이름이 정확히 같아야 일치로 보기 때문에 보류됐고, 담당 직원이 같은 회사의 입금으로 확인해 승인했습니다. 10월 21일부터는 고객사 소명이 없는 보류 건은 무엇으로 명의를 확인했는지 10자 이상 적어야 승인되고, 그 근거가 승인 기록에 남습니다.」',
     'H-12', '10-21 실측 = 보류 14:21:44 → 승인 14:22:31 eunji_admin · 소명 없음 · 메모 「무통장입금」 · 입금자 「임상옥 지삼유통」 · ⚠ 임상옥 = 등록 대표자인지 확인 전 · 보완 문장 = 10-21 배포 S4(deposit-approve.ts confirmBasis)'),
    ('Q-46', '2.3', '우회 충전 · 수동 충전 · 이미 충전됨 · 베이컨',
     '9월 18일 건은 「이미 충전됨」으로 반려됐습니다. 보류 중에 다른 길로 충전한 것 아닙니까?',
     '「맞습니다. 보류 7분 뒤 담당자가 같은 금액을 관리자 수동 충전으로 먼저 넣었습니다. 수동 충전 화면이 입금 신청을 보지 않았기 때문입니다. 이 사례를 계기로 10월 5일부터 명의 확인을 기다리는 입금이 있으면 수동 충전을 시스템이 거부합니다.」',
     'H-12', '문서 2.3 5절에 같은 문장 · 수동 충전 거부 시험 화면(10/05) 별첨'),
    ('Q-47', '2.3', '시험 계정 보류 · 홍길동 · 대표자 명의',
     '테스트계정 보류 건 두 개는 무엇입니까?',
     '「당사 시험 계정으로 보류 · 소명 · 처리 기능을 점검한 건입니다(8월 29일 · 10월 5일). 둘 다 반려했습니다.」',
     'H-12', '10-21 실측'),
    ('Q-48', '2.3 · 4.1', '입금자 이름 원문 · 감사 기록 개인정보',
     '처리 기록에 입금자 이름이 원문으로 남아 있습니다.',
     '「입금자명은 명의 대조에 원문이 필요한 값입니다. 로그 정책에서 저장할 때 가린다고 한 것은 전화번호와 4자리 이상 숫자열이고, 감사 기록은 대표 등급만 조회합니다.」',
     'H-12 · H-1', '문서 4.1-4.3 1절 · 5-2 접근권한 · 실사 화면에는 첫 글자만 나오게 조회문을 고쳤다(D-18)'),
    ('Q-49', '3.1 ③ · ④', '사설 원격 접속망 · Tailscale · 출장 · 우회 경로',
     '관리 접속은 SSH 하나라고 하셨는데, 출장 때 쓴 사설 원격 접속망은 지금 어떻습니까? 켜고 끈 기록이 있습니까?',
     '「출장 중(10월 8일~11일)에는 대표 노트북 한 대만 이 길로 관리 접속했고, 귀국 뒤에는 쓴 적이 없습니다. 해제가 늦어진 것을 10월 21일 점검에서 찾아 두 서버 모두 끄고, 출장 때 넣은 방화벽 줄을 지우고, 운영 기록 대장에 켠 날과 끈 날을 남겼습니다.」',
     'H-21g · G-12f · H-21c · H-16', '10-21 실측 = 두 서버 접속 기록에서 100.x 출발지 = 100.116.151.70(대표 노트북) 하나 · 10/08~10/11 23:07 · 그 뒤 0 · 10-21 두 서버 disable --now(inactive · disabled) · .62 ufw trip-tailscale-ssh 줄 삭제 10:13:52 · 대장 4건 등록(서수란 10:29~10:32) · 대표 확인 10:37 = H-16 · 10/08 13:31 crash= .62 서버 장애 재부팅(SERVERS.md ★1008)'),
    ('Q-50', '3.2', '직원 계정 현황 · OTP 등록 · 초기 비밀번호',
     '사용 중인 직원 계정이 모두 인증 앱을 등록하고 초기 비밀번호를 바꿨는지 보여 주십시오.',
     '「사용 중인 직원 계정은 대표 1 · 지원팀장 1 · 지원팀원 2 모두 네 개이고, 모두 인증 앱을 등록했고 비밀번호 변경이 남은 계정이 없습니다. 사용 중지된 계정은 두 개입니다.」',
     'H-15 첫 표', '10-21 실측 = ceo · eunji_admin · jihyun_admin · suran 사용 · OTP t · 변경 필요 f / admin · psypsy 중지'),
    ('Q-51', '3.2 ③', '등급 변경 사유 · 조정 · 3등급',
     '등급 변경 사유가 「조정」뿐입니다. 무슨 조정입니까?',
     '「8월 27일 직원 등급을 대표 · 지원팀장 · 지원팀원 세 단계로 나누면서, 지원팀원 두 명을 대표 등급에서 지원팀원 등급으로 낮춘 것입니다. 하는 일에 필요한 만큼만 권한을 두었습니다.」',
     'H-15 둘째 표', '10-21 실측 = 8/27 23:17 · 23:18 eunji_admin · jihyun_admin super → support · 처리자 ceo · 3등급 도입일 = 8/27(원장 §4-I)'),
    ('Q-52', '3.2', 'psypsy · 박성용 · 심사 계정 · 1분 중지',
     'psypsy 계정은 심사 때문에 만들었다가 1분 만에 중지했습니다. 무엇입니까?',
     '「전송자격인증 심사 준비 중에 직원 계정을 잘못 만들어 1분 뒤 사용 중지했습니다. 한 번도 로그인하지 않았고 인증 앱도 등록하지 않아 쓸 수 없던 계정입니다. 중지 사유 칸에는 그 계정을 만든 경위를 적었습니다.」',
     'H-15', 'Harold 10-21 = 잘못 생성해 정지 · 10-21 실측 = 8/27 23:44:17 발급(사유 「신규발급」 · 지원팀원) → 23:45:45 중지(사유 「전송자격인증제 심사위한 계정생성」) · 로그인 0번 · OTP 없음 · 비밀번호 변경 필요 t = 쓸 수 없는 계정 · 사유 문장이 보여주기용으로 읽힐 수 있다'),
    ('Q-53', '3.2 ③', '직원 계정 발급 기록 · 8월 27일 이전 · 생성일',
     '대표 · 지원팀원 계정은 발급 기록이 없습니다. 누가 언제 승인해 만들었습니까?',
     '「직원 계정 발급 기록은 8월 27일부터 남깁니다. 그 전에 만든 계정(대표 · 지원팀원 두 명)은 생성일로 확인합니다. 그 뒤 발급은 처리자 · 사유와 함께 기록됩니다.」',
     'H-15', '10-21 실측 = admin_account_created 첫 기록 8/27 · ceo 5/6 · eunji · jihyun 5/11 생성'),
    ('Q-54', '3.4 · 3.2', 'OTP 재등록 · 대표 OTP 두 번',
     '대표 계정 OTP 등록 기록이 두 번(5월 11일 · 8월 25일)입니다.',
     '「8월 25일 인증 앱을 쓰던 기기를 바꿔서 다시 등록했습니다.」',
     'H-15 둘째 표', 'Harold 10-21 = 기기 교체 · 10-21 실측 = ceo totp_enrolled 5/11 · 8/25'),
    ('Q-55', '공통', '시각 9시간 · UTC · 표기 차이',
     '어떤 표는 시각이 9시간 차이 납니다.',
     '「일부 표는 세계 표준시(UTC)로 저장돼 있어 그대로 찍으면 9시간 빠르게 보입니다. 같은 순간이고, 오늘 보여 드리는 조회는 한국 시각으로 맞춰 두었습니다.」',
     'H-15 · H-8b', '10-21 실측 = super_admins · callback_numbers 가 UTC 저장(감사 기록과 소수점까지 같고 9시간 차이)'),
    ('Q-56', '3.3 · 4.3', '대장 무효 처리 · 13건 · 작성자 착오',
     '운영 기록 대장에 무효 처리된 기록이 많습니다. 지운 것입니까?',
     '「대장을 처음 열 때(10월 3일) 대표 이름으로 올렸는데 점검 담당은 지원팀장이라, 작성자 착오로 무효 처리했습니다. 지우지 않았고 원래 기록과 무효 처리 기록이 그대로 남습니다. 같은 내용을 지원팀장이 다시 올리고 대표가 확인했습니다. 10월 21일 한 건도 같은 이유입니다.」',
     'H-16', '10-21 실측 = 10/03 대표 13건 21:40 무효 · 10/04 서수란 16건 → 10/05 Harold 확인 · 10/21 10:20 대표 1건 10:21 무효'),
    ('Q-57', '3.3', '접근권한 점검 정정 · 두 줄',
     '3분기 접근권한 점검이 두 번 있습니다.',
     '「대상 시스템을 「한줄로」로 잘못 골라 「공통」으로 정정해 다시 올렸습니다. 둘 다 확인을 마친 뒤라 두 줄이 남았고, 기록은 고치지 않습니다.」',
     'H-16', '10-21 실측 = 10/04 22:59 hanjul · 23:01 common(정정)'),
    ('Q-58', '3.4 ①', '인증 요청과 성공 차이 · 미완료',
     '인증 요청보다 성공이 적습니다(9월 56 · 52, 10월 114 · 109). 나머지는 무엇입니까?',
     '「인증번호를 받고 입력하지 않았거나 다시 받은 건입니다. 접속 세션은 인증번호가 맞을 때만 발급되므로 인증 없이 들어간 접속은 없습니다.」',
     'H-13', '문서 3.4 3절 4 · 10-21 실측 = 실패 9월 1 · 10월 0'),
    ('Q-59', '3.4 ①', '매 로그인 인증 · 기기 기억 없음',
     '로그인할 때마다 인증한다는 것을 보여 주십시오.',
     '「오늘 오전 같은 계정이 2분 동안 네 번 로그인했는데, 네 번 모두 인증번호를 요청하고 성공해야 들어갔습니다.」',
     'H-13 셋째 표', '10-21 실측 = psy5868 10:26 · 10:27 · 10:27 · 10:28 요청 → 성공'),
    ('Q-60', '3.4 ④', '계정 잠금 · 5회 실패 · 잠금 사례',
     '인증번호 5회 실패 잠금이 실제로 동작한 사례가 있습니까?',
     '「실제 잠금은 아직 없었습니다. 인증번호 실패는 9월 1건뿐입니다.」',
     'H-13', '10-21 실측 = 잠김 0건'),
    ('Q-61', '3.4 ③', '번호 수동 변경 · 본인확인 뒤 변경 · gwchae',
     '본인확인을 마친 계정(gwchae)의 번호를 10월 13일 관리자가 바꿨습니다. 본인확인 번호를 손으로 바꿀 수 있는 것 아닙니까?',
     '「본인확인 번호는 그대로이고, 10월 13일 기록은 같은 번호로 다시 저장된 것입니다(변경 전후 같음). 본인확인을 마친 다섯 계정 모두 지금 수신 번호가 본인확인 번호와 같습니다.」',
     'H-13b · H-7b', '10-21 실측 = 10/13 16:58 ceo → gwchae mfa_phone_changed · gwchae 본인확인 10/07 08:51 · 문서 3.4 2절 「손으로 바꿀 수 없다」'),
    ('Q-62', '3.4 ③', '대상 계정 빈칸 · 번호 변경',
     '10월 20일 번호 변경 한 건은 대상 계정이 비어 있습니다.',
     '「이폴리움이 선불 금액을 따로 관리하겠다고 요청해 고객사를 둘로 나눴습니다. 그때 기존 계정을 새 고객사로 옮기면서 옛 계정을 지우고 같은 아이디로 새로 발급했습니다. 옛 계정이 지금 없어서 대상 칸이 비어 보입니다. 그때는 계정 삭제가 감사 기록에 남지 않았고, 점검에서 찾아 10월 21일부터 계정 삭제도 감사 기록에 남도록 보완했습니다.」',
     'H-13b · H-6', '10-21 실측 = 옛 efolium b7d2f5a9 · 15:11:42 수신거부 248건 내보내기 · 15:12:40 번호 저장(전후 빈 값) · 그 뒤 계정 표에 없음 · 15:13:00 새 efolium ea95101c 발급 · 고객사 분리 = 이폴리움 요청(선불 금액 따로 관리 · 10-21 Harold) · 삭제 기록 없음 = D-24 → 10-21 배포 S1(user_account_deleted)'),
    ('Q-63', '3.4 ③', 'psy5868 · 번호 일곱 번 변경 · 시험 계정',
     'psy5868 계정 번호를 9월에 일곱 번 바꿨습니다.',
     '「psy5868 은 당사 내부 관리 계정입니다. 카카오 알림톡 템플릿이 이 계정에만 저장돼 있어서, 대표와 영업팀장이 번갈아 시험할 때 인증번호를 받을 번호를 두 사람 번호로 바꿔 저장했습니다. 두 번호만 오갔고 모두 당사 직원 번호입니다.」',
     'H-13 넷째 표', '10-21 실측 = 9/12 · 9/19 두 번 · 9/20 · 9/21 두 번 · 9/25 · 처리자 ceo · 이유 = Harold 10-21(템플릿이 이 계정에만 있어 대표 · 영업팀장 번호 끝 8517 · 5868 만 오감)'),
    ('Q-64', '3.5 ① · 6절', '발신 인증 기록 · 발신번호 · 수신 번호',
     '발신 인증 기록에 무엇이 남습니까?',
     '「계정 · 발신번호 · 인증번호를 받은 번호(가운데 가림) · 접속 IP · 일시가 요청과 성공 때마다 남습니다.」',
     'H-14 둘째 표', '10-21 실측 = phone_masked · callback_number'),
    ('Q-65', '3.5 ④', '발신번호 * · 계정 단위 인증 · 회신번호',
     '발신번호가 「*」로 인증된 기록이 있습니다. 모든 번호가 인증되는 것입니까?',
     '「아닙니다. 고객마다 회신번호가 다른 발송처럼 발신번호를 하나로 정할 수 없는 발송은 번호 대신 계정 단위로 인증을 받습니다. 이 인증은 그런 발송에만 쓰이고, 특정 번호의 인증으로 넓혀지지 않습니다. 번호를 모른다고 인증을 건너뛰지 않기 위한 장치입니다.」',
     'H-14 둘째 표', '근거 = utils/sender-auth.ts:70-78 SENDER_AUTH_ANY_CALLBACK · 10-21 실측 = gwchae 10/19 10:52'),
    ('Q-66', '3.5 ④ · 5절', '재인증 · 8시간 · 발신번호 변경',
     '8시간이 지나거나 발신번호가 바뀌면 다시 인증한다는 것을 보여 주십시오.',
     '「같은 직원이 10월 19일 10시 22분에 대표번호로 인증하고, 10시 31분에 다른 발신번호로 바꾸자 다시 인증했습니다. 또 다른 직원은 10월 19일 14시 53분 인증 뒤 다음 날 13시 59분에 다시 인증했습니다(8시간 경과).」',
     'H-14 둘째 표', '10-21 실측 = sgbaek 10/19 10:22(18008125) → 10:31(01082336860) · suran 10/19 14:53 → 10/20 13:59'),
    ('Q-67', '4.1', '화면 접속 기록 · 페이지 접속',
     '이용자의 화면 접속 기록을 보여 주십시오.',
     '「로그인한 이용자가 어느 화면에 들어갔는지 계정 · 화면 주소 · 접속 IP · 일시로 남깁니다. 같은 사람이 같은 화면을 1분 안에 다시 열면 한 줄로 묶습니다. 9월 7,758건 · 10월 8,122건입니다.」',
     'H-4', '10-21 실측 · 2026-09-12부터'),
    ('Q-68', '4.2', '개인정보 조회 기록 · 시작일 · 0건 조회',
     '개인정보 조회 기록은 언제부터입니까? 결과가 없는 조회도 남습니까?',
     '「조회 · 등록 · 수정 기록은 9월 12일부터, 내보내기 · 삭제 기록은 그 전부터 남깁니다. 결과가 0건인 조회도 화면 이름과 함께 기록됩니다.」',
     'H-5', '10-21 실측 = gwchae 10/21 09:05 count 0 · 내보내기 첫 기록 8/20 · 삭제 2026-02'),
    ('Q-69', '4.2', '고객 선택삭제 · 2월 281건 · 고객 DB 삭제',
     '2월 고객 선택삭제 281건은 무엇입니까?',
     '「고객사가 자기 고객 DB에서 고객을 골라 지운 기록입니다. 지우는 요청 한 번이 한 줄로 남습니다.」',
     'H-5 첫 표', '10-21 실측'),
    ('Q-70', '4.1 · 3.4', '로그인 자동 차단 · 5회 실패 · 30분',
     '로그인 실패가 쌓이면 어떻게 됩니까?',
     '「같은 출발지 · 같은 아이디로 10분 안에 5번 틀리면 30분 동안 자동으로 막고 화면에 안내합니다. 30분이 지나거나 관리자가 확인해 풀면 해제되고, 차단 · 해제 시각과 사유가 남습니다.」',
     'H-3 둘째 표', '10-21 실측 = r0029 10/21 09:46 차단 → 10:16 만료'),
    ('Q-71', '4.1 · 2.2', 'hajulai · 일본 IP 차단 · 아이디 오타',
     '10월 10일 일본 주소에서 로그인 실패 차단이 세 번 있습니다.',
     '「대표가 출장 중(예외 승인 기간)에 시연 계정 아이디를 잘못 입력해(hanjulai 를 hajulai 로) 5번 실패해 막힌 것이고, 대표가 확인해 풀었습니다. 그런 아이디의 계정은 없습니다.」',
     'H-3 둘째 표', '10-21 실측 = 219.164.64.138 · 19:37 · 19:39 · 19:39 · admin_unblock'),
    ('Q-72', '4.1 · 3.4', '직원 계정 반복 차단 · eunji_admin · 8월 27일',
     '8월 27일 직원 계정이 10분 사이 아홉 번 막혔다가 풀렸습니다.',
     '「기록상 사무실 주소에서 로그인이 반복해 막혔고, 관리자가 본인 확인 뒤 풀었습니다. 차단과 해제가 모두 처리자 · 시각과 함께 기록에 남아 있습니다.」',
     'H-3 둘째 표', '10-21 실측 = 180.226.236.94(사무실) · 15:14~15:26 · 전부 admin_unblock'),
    ('Q-73', '5.1', '고객사 해지 · 해지 사유 · 해지 고지',
     '고객사를 해지할 때 사유와 고객 고지가 기록됩니까?',
     '「10월 21일부터 해지는 회사 목록의 해지 버튼 한 길로만 됩니다. 사유(계약 종료 · 불법스팸 신고 · 관리자 조치)를 반드시 고르고, 상태 변경 · 소속 계정 잠금 · 접속 끊기 · 해지 기록이 한 번에 처리됩니다. 고객사 관리자에게 사유 안내 문자가 나가고 그 결과도 따로 기록됩니다. 그 전 해지 기록 2건은 사유 칸이 없던 때 기록입니다.」',
     'H-18b · H-18', '10-21 실측 = 자연인 9/07(잠금 3 · 세션 2) · 엔터식스 10/06(잠금 2 · 세션 0)· 사유 · 고지 칸 없음 · 10-21 배포 S3(utils/company-terminate.ts) · 실측 10-21 13:39 테스트주식회사 해지(사유 · 잠금 2 · 문자 1 · 대표 수신) = H-18b 로 보여 준다'),
    ('Q-74', '5.1', '스팸 신고 조치 사례 · 이용 정지 · 계정 제한',
     '스팸 신고로 이용을 제한하거나 정지한 사례가 있습니까?',
     '「불법스팸 발송이 한 건도 확인되지 않아 실제 조치 사례는 없습니다. 조치 절차는 10월 3일 당사 시험 계정으로 한 번 돌려 봤고, 상태 변경 · 접속 종료 · 문자 고지가 함께 기록됐습니다.」',
     'H-18', '10-21 실측 = invito01 active → dormant · notified true · killedSessions 1'),
    ('Q-75', '5.2', '금칙어 차단 규칙 · 탐지 · 차단 전환',
     '차단 체계라고 하셨는데 지금 차단 모드 규칙이 있습니까?',
     '「새 규칙은 탐지로 시작해 실제 발송 문안에서 오탐이 없는지 본 뒤 차단으로 바꿉니다. 불법스팸 규칙은 8월 27일부터 두 달 탐지로 돌려 오탐이 없어 10월 21일 차단으로 바꿨습니다.」',
     'H-19', '10-21 실측 = 불법스팸(대출 + 리딩방) 13:53:44 detect → block(ceo · 전환 기록 spam_block_rule_mode) · 탐지 기록 0 · 시험 규칙은 13:53:46 잘못 눌러 13:54:10 detect 로 되돌림(물으면 「시험 규칙을 잘못 눌러 바로 되돌렸다 · 기록이 그대로 남는다」) · 문서 1.3 5-1 「탐지로 시작 · 오탐 확인 뒤 차단 전환」'),
    ('Q-76', '5.2', '금칙어 차단 실증 · 차단 기록',
     '금칙어로 실제 발송이 막힌 기록을 보여 주십시오.',
     '「10월 3일 시험 규칙을 차단으로 돌려 당사 시험 계정으로 보낸 문자가 요금 차감 전에 막혔고, 발송자에게 차단 안내가 나갔습니다. 그 기록입니다(문안 숫자는 가려 저장).」',
     'H-19 셋째 · 넷째 표', '10-21 실측 = 10/03 02:45 UTC(한국 11:45) · direct_core · block · 1건'),
    ('Q-77', '5.2', 'KISA 차단정보 · 외부 출처 규칙',
     'KISA 등 외부기관 차단정보를 규칙으로 등록한 것이 있습니까?',
     '「KISA 악성문자 사전차단 데이터(SHARE)를 수시로 받아 당사가 자체 개발한 링크가드 관제센터에 출처 KISA 로 적재하고, 발송 엔진이 문자 속 URL 을 발송 전에 이 목록과 대조해 막습니다. 한줄로 금칙어 화면은 당사가 직접 넣는 문구 규칙이고, KISA 악성 URL 목록은 링크가드에 있습니다.」',
     'H-19 · 링크가드 관제센터(URL 정책 · 판정 이력 · 대표 화면)', '10-21 Harold 관제센터 화면 = KISA 마지막 적재 10/21 08:47 · URL 정책 차단 목록 출처 KISA SHARE(kisa-sync · seq 176017 등) · 판정 이력 목록 일치 차단(9/19 ~ 10/2 bito-gateway) · 한줄로 금칙어 규칙 출처는 전부 internal(H-19)'),
    ('Q-78', '4.1 ⑤ · 4.2', '앱 로그 1년 · 로그 회전 · 400일',
     '서버 응용 프로그램 로그를 1년 이상 보관하는 설정을 보여 주십시오.',
     '「응용 프로그램 로그는 매일 회전해 400개(400일)를 압축 보관합니다. 9월 12일 설정 뒤 파일이 그대로 쌓여 있고, 로그 폴더는 실행 계정만 쓸 수 있습니다.」',
     'H-22a', '10-21 실측'),
    ('Q-79', '4.1 ⑤', '백업 · 암호화 · 외부 전송 · 경보',
     '백업이 실제로 매일 도는지 보여 주십시오.',
     '「매일 새벽 3시에 두 데이터베이스를 암호화해 백업하고 외부 보관 서버로 보냅니다. 오늘도 3시 13분에 성공했고, 매일 8시 30분 감시가 최근 성공 여부를 확인해 문제가 있으면 담당자 휴대폰으로 알립니다.」',
     'H-23', '10-21 실측 = LAST_SUCCESS 10-21 03:13 · 감시 OK 매일'),
    ('Q-80', '4.2', 'DB 로그 오류 · 컬럼 없음 · 장애 로그',
     'DB 로그에 오류가 반복해서 찍혀 있습니다. 점검하고 조치했습니까?',
     '「점검에서 두 기능의 조회 오류를 찾아 10월 21일 고쳤습니다. 그 뒤로 몇 분마다 반복되던 오류는 없어졌고, 한 기능은 6시간에 한 번 확인 기록만 남습니다. 두 오류 모두 기능 결과에는 영향이 없었습니다.」',
     'H-22c', '10-21 실측 = cdp_push_subscriptions.is_active · kakao_templates.imc_template_status 컬럼 없음 반복 · 10-21 배포 S5 · 재기동 12:50:14 · 실측(Harold 실행) = is_active 마지막 12:48:21(재기동 전) · 재기동 뒤 15분 0건 · imc 12:51:32 1건(재기동 직후 한 번 · 6시간 기억) · 웹푸시 구독 표 0행(판정 결과 전후 같음)'),
    ('Q-81', '3.1 · 4.1 ⑤', '외부 열린 DB · 23388 · 결제 통계 수신',
     '컨테이너 목록에 외부로 열린 MySQL(23388)이 있습니다. 데이터베이스는 내부 전용이라고 하지 않았습니까?',
     '「발송 중계사가 결제 통계를 넣는 수신 전용 데이터베이스입니다. 서비스 데이터베이스(PostgreSQL · MySQL · Redis)는 서버 내부 주소에만 열려 있고, 이 수신 포트는 컨테이너 전용 방화벽에서 송신 서버만 허용하고 그 밖은 막습니다(방화벽 운영정보 4절).」',
     'H-21e · H-21h', '10-21 실측 = DOCKER-USER 23388 허용 5곳 뒤 DROP = 문서 3.1 4절과 같음'),
    ('Q-82', '4.1 ⑤', '백업 비밀번호 경고 · mysqldump',
     '백업 기록에 「명령줄 비밀번호」 경고가 있습니다.',
     '「백업은 서버 안에서 관리자 계정으로만 실행되고, 서버 원격 접속은 지정된 두 곳에서만 됩니다(방화벽). 그 문구는 데이터베이스 도구가 비밀번호를 명령으로 넘길 때마다 띄우는 안내이고, 백업 파일은 공개키로 암호화돼 서버에 복호화 키가 없습니다.」',
     'H-23', '10-21 실측'),
    ('Q-83', '4.1 · 3.4', '차단 해제한 사람 · 관리자 해제',
     '로그인 차단을 관리자가 풀었다는데 누가 풀었습니까?',
     '「해제한 관리자 계정이 함께 남습니다. 8월 27일 해제는 당시 공용 계정(그날 밤 사용 중지)이, 그 뒤는 지원팀장과 대표가 했습니다.」',
     'H-3 둘째 표', '10-21 실측 = 8/27 admin · 9/11 · 10/20 suran · 8/27 밤 · 8/28 · 10/10 ceo · ⚠ 10/10 19:39:06 한 줄은 해제한 사람이 비어 있다'),
    ('Q-84', '4.2 · 4.1 ⑤', '시스템 저널 · 1년 보관 · 8월 15일 · 파일 개수 한도',
     '시스템 저널을 1년 보관한다고 했는데 8월 15일부터만 있습니다.',
     '「시스템 저널이 파일 개수 기본 한도(100개)에 걸려 8월 중순 이전분이 자동으로 지워진 것을 10월 21일 점검에서 찾아 한도를 고쳤습니다. 이용자 로그와 감사 기록은 데이터베이스에 따로 보관돼 지워지지 않았습니다(2월 11일부터 그대로). 경위는 운영 기록 대장에 남겼습니다.」',
     'H-22b · H-1', '10-21 실측 = 저널 폴더 1개 · 파일 100개 1.4G · 첫 기록 08-15 · 한도 변경 10-21 · 대장 로그 점검 10-21 등록 = H-16'),
    ('Q-85', '3.1 ①', '방화벽 현재 정책 · 원격 관리 두 곳',
     '지금 방화벽 정책을 보여 주십시오.',
     '「들어오는 연결은 기본 차단이고, 웹(80 · 443)만 열려 있습니다. 원격 관리(22)는 대표 · 사무실 두 곳에서만 되고, 내부 응용 포트(3000 · 9001~9011)는 명시적으로 막았습니다.」',
     'H-21a', '10-21 실측 · 출장용 줄 삭제 뒤'),
    ('Q-86', '4.1', '엔진 발송 이력 · 보관 범위 · 사설발송',
     '사설발송 엔진의 발송 이력을 보여 주십시오.',
     '「엔진 운영을 시작한 6월 15일부터 모든 발송 요청이 남아 있습니다. 요청 시각 · 발송 계정 · 수신번호(가림) · 발신번호 · 유형 · 상태 · 결과 코드 · 결과 시각이 건마다 남습니다.」',
     'G-11', '10-21 실측 = 2026-06-15 12:38 ~ 10-21 11:03'),
    ('Q-87', '4.2', 'lock timeout · DB 오류 · 엔진 DB 로그',
     '엔진 DB 로그에 「lock timeout」 오류가 있습니다.',
     '「같은 연동 계정 정보를 두 작업이 동시에 고치려다, 뒤의 작업이 정해 둔 대기 한도를 넘어 스스로 취소된 기록입니다. 오래 기다리며 다른 처리를 막지 않게 둔 안전장치이고, 취소된 작업은 아무것도 바꾸지 않았습니다.」',
     'G-13c', '10-21 실측 = 10:02:03 agent_account FOR UPDATE 1건 · 문장이 취소되면 그 트랜잭션은 되돌려진다(PostgreSQL)'),
    ('Q-88', '3.1 ④', '서버 접속 수집 · 지금도 수집 · 게이트웨이 서버 명령',
     '게이트웨이 서버 접속 · 명령 수집이 지금도 동작합니까?',
     '「동작합니다. 오늘 아침 서버 로그인 · 명령 기록이 들어와 있고, 수집기가 몇 분마다 살아 있다는 신호를 남깁니다.」',
     'G-6', '10-21 실측 = SERVER_LOGIN 10:04 · SERVER_COMMAND 11:09 · HEARTBEAT 11:04'),
    ('Q-89', '2.2 ①', '사설 주소 허용 IP · 10.0 · eternal',
     '허용 IP에 사설 주소(10.0.x)가 있습니다.',
     '「고객사 클라우드 서버가 알려 준 내부 주소를 공인 주소와 함께 적은 것입니다. 바깥에서 들어오는 접속은 공인 주소로 판정하므로 사설 주소로는 들어올 수 없습니다.」',
     'G-5 · G-3', '10-21 실측 = eternal01 10.0.131.103 · eternal02 10.0.12.170 · ⚠ 정리 여부 = D-33'),
    ('Q-90', '2.2 ③', '허용 IP 사유 · 실제 기록 · 만료일',
     '허용 IP를 바꿀 때 사유가 필수라는데 사유가 적힌 기록이 있습니까?',
     '「사유 칸을 넣은 10월 3일 뒤로는 허용 IP를 바꾼 일이 없어서 사유가 적힌 줄은 아직 없습니다. 바꾸면 사유 없이는 저장되지 않고, 만료일을 둘 수 있습니다(관리 화면).」',
     'G-3 · G-5', '10-21 실측 = 마지막 허용 IP 변경 10/03 09:13(api-rabd-api-01 · 사유 칸 배포 09:29 전)'),
    ('Q-91', '3.4 ⑥', '설치 묶음 반복 발급 · itensms03',
     '같은 계정에 설치 묶음을 여러 번 발급했습니다.',
     '「고객사 서버 설치를 도우면서 여러 번 다시 발급한 기록이며, 발급마다 처리자와 일시가 남습니다.」',
     'G-7', '10-21 실측 = itensms03 9/04 ~ 9/09 12번'),
    ('Q-92', '4.3 · 3.1 ④', '관리 사이트 로그인 수 · 점검 자료',
     '엔진 관리 사이트 로그인이 월별로 몇 번입니까?',
     '「OTP까지 마친 로그인이 9월 88번 · 10월 49번이고, 실패는 9월 10번 · 10월 6번입니다.」',
     'G-6 둘째 표 · G-10', '10-21 실측 · 9월 점검 기록(대장) 「관리 사이트 로그인 88」과 같음'),
]

# (번호, 심사 항목, 핵심 검색어, 결함, 처리 = 수정 · 운영 조치 · 반박, 처방 또는 반박 근거, 상태, 관련 질답)
# 목표 = 트집거리 0(★2026-10-21 Harold). 수정은 설계안 → 동의 → 구현 → Codex → 배포. 운영 조치 작성일은 오늘(소급 금지).
DEFECTS = [
    ('D-01', '2.1', '발급자 빈칸 · 스크립트',
     '서버 관리 스크립트로 만든 시연 · 시험 회사 계정 발급 기록에 발급자가 없다(경로 표시는 운영자 발급)',
     '수정 + 반박', '뿌리 = 서버 관리 스크립트가 감사 기록에 실행자를 넘기지 않는다(D-02 와 같은 뿌리) · 수정 = 스크립트 실행자를 감사 기록에 남김 / 반박 = 당사 내부 시연 · 시험 회사뿐 · 고객사 발급은 전부 발급자가 남음(H-6) · 지난 기록은 바꾸지 않는다',
     '배포 완료 10-21(Codex 적대 3R approve) · 실측 대기 · ⚠ 지난 스크립트 실행자 = Harold', 'Q-03'),
    ('D-02', '2.1', '시연 회사 삭제 · 삭제 기록',
     '은퇴한 시연 회사를 정리 스크립트로 지울 때 감사 기록이 남지 않았다',
     '수정 + 반박', '10-21 실측 = 백업 파일 administrator · root 홈 둘 다 없음 → 삭제 증거가 남아 있지 않다 · 수정 = 정리 스크립트가 삭제를 감사 기록에 남김(D-01 과 한 묶음) / 반박 = 당사 내부 시연 회사 · 시연 표식 회사만 대상 · 감사 기록 표는 회사 삭제로 지워지지 않는다(H-6c 15줄 남음)',
     '배포 완료 10-21(Codex 적대 3R approve) · 실측 대기', 'Q-06'),
    ('D-03', '2.1', '발신번호 삭제 · 배정 해제 · 삭제 기록',
     '발신번호 삭제(관리자 · 고객사 두 경로)와 배정 해제(두 자리)에 감사 기록이 없다',
     '수정', '감사 기록 컨트롤타워 하나를 네 자리에 연결(처리자 · 대상 번호 · 고객사 · 일시) · 같은 패턴(2.1 수정 · 삭제 경로 전부) 전수 grep 뒤 함께 · 지난 삭제(시험 번호 등)는 「오늘 전에는 기록이 없었다」로 반박',
     '배포 완료 10-21 · 실측 완료(13:15 테스트계정 02-1234-5678 등록 → 삭제 두 줄 · H-8c)', 'Q-26 · Q-23 · Q-24'),
    ('D-04', '2.2 · 4.3', '10월 7일 · 국내 대역 1개 · 오차단',
     '10/07 15:47 국내 대역 표가 2,597 → 1개로 교체돼 국내 이용자 4계정이 9분간 국외로 막혔다 · 16:07 되돌림은 감사 기록이 없다',
     '운영 조치 + 반박', '수정은 10/07 당일 끝남(일괄 교체 경로 · 화면 칸 삭제 · status/OPS.md 2-2-F · 지금 소스에 경로 없음) / 운영 조치 = 운영 기록 대장에 이상 · 원인 · 조치 · 재확인을 오늘 작성일로 기록(SQL 복구라 감사 기록이 없는 부분 보완) / 반박 = 경위 순서(H-9b)',
     '닫힘(10-21 대장 로그 점검 등록 · H-16)', 'Q-32 · Q-34'),
    ('D-05', '2.1 ③', '테스트계정 · 유선 상한 1',
     '10/05 시험으로 테스트계정 유선 상한을 1로 바꾼 뒤 되돌린 기록이 없다',
     '운영 조치', '지금 값을 확인하고 1이면 화면에서 되돌린다(화면 경로는 기록이 남는다)',
     '닫힘(10-21 13:15:20 대표가 화면에서 유선 상한 1 → 5 · 회선 정책 변경 기록 남음 · 지금 5)', 'Q-27'),
    ('D-06', '3.2', 'admin 계정 · 공용 계정',
     '8/20 발신번호 심사자가 admin 계정 · 누가 쓰는 계정인지에 따라 공용 계정 지적',
     '반박', 'Harold 10-21 = 공용 계정이라 사용 중지하고 담당자마다 개인 계정 발급(문서 3.2 1-6절과 같음) · 10-21 실측 = 중지 2026-08-27 23:18 → 8/20 심사 = 정리 전 기록',
     '닫힘', 'Q-28'),
    ('D-07', '2.1 ②', '중복확인값 · 중복가입',
     '문서에 「중복가입 확인값을 해시로 기록」 · 실제로 값이 차 있는지 미확인',
     '반박', '10-21 실측 = 인증 완료 6건 전부 기록됨(t) · 문서와 시스템 일치',
     '닫힘', 'Q-11'),
    ('D-08', '2.1', '4월 · 7월 계정 몰림',
     '4월 153 · 7월 53 계정이 한꺼번에 생겼다',
     '반박', '관리자 수 = 고객사 수(H-6b) + 그때 한 작업',
     '답변 확정(Harold 10-21 = 4월 레거시 이전 · 7월 에이전트 통계)', 'Q-05'),
    ('D-09', '2.1 ⑤', '반려 사유 · 시험 입력',
     '시험 고객사 반려 사유가 「ㅁㅇㄴㅇㅁㄴㅇ」',
     '반박', '회선 상한 시험 신청 · 기록은 고치지 않는다',
     '답변 확정', 'Q-22'),
    ('D-10', '2.1 ⑤', '발신번호 등록 시각 · UTC',
     '발신번호 표의 등록 시각이 9시간 빠르게 나온다(UTC 저장)',
     '반박', '표기 차이 · 같은 순간(소수점까지 심사 시각과 같음)',
     '답변 확정', 'Q-24'),
    ('D-11', '2.1 ②', '고객사 본인확인 · 10월 26일',
     '고객사 계정 본인확인 기록이 아직 없다',
     '반박', '문서 2.1 3절 「10/06 운영 개시 · 10/26 전 계정」 · 로그인 화면 사전 고지',
     '답변 확정 · 10/26 운영 조치 = 본인확인 · 다중인증 · 발신 인증 명단을 전 계정(*)으로', 'Q-10'),
    ('D-12', '2.2 · 3.4', '설정 파일 중복 줄 · GEO_BLOCK_ENFORCE_FROM',
     '운영 설정 파일에 국외 차단 시행일 줄이 두 번 있다(값은 같다 · 2026-10-01)',
     '운영 조치', '한 줄을 지운다 · 값이 같아 동작은 바뀌지 않는다 · 실사 중 H-24 를 띄우면 같은 줄 두 개가 보인다',
     '닫힘(10-21 대표 · 95번째 중복 줄 삭제 · 백업 .env.bak-1021-geo · 값 같아 동작 불변)', '-'),
    ('D-13', '2.2 ④', 'VPN 예외 · 시세이도 · shadmin',
     '예외 사유가 「시세이도 본사 VPN」 · 우회접속(VPN)을 허용한 것으로 읽힐 수 있다',
     '반박', '고객사 본사가 의무로 쓰는 사내망의 고정 출구 주소 하나(/32)만, 그 계정 하나에만, 사유 · 승인자를 적어 허용 · 문서 2.2 6절 「업무상 반드시 거치는 사내 보안망만」과 같음',
     '답변 확정', 'Q-35'),
    ('D-14', '2.2 ③', '대표 출장 예외 · 전역 범위 · 유효 표시',
     '대표 출장 예외가 계정이 아니라 전역(global) 범위 · 만료(10/12)가 지났는데 대장에 「유효 t」',
     '반박', '유효 = 회수 여부 · 만료는 판정 때마다 본다(utils/geo-access.ts:246) · 슈퍼관리자 계정은 고객 계정 칸에 넣을 수 없어 그 주소 하나(/32)만 4일 기한으로 허용 · 슈퍼관리자 계정 범위는 추가 과제',
     '답변 확정', 'Q-36 · Q-37'),
    ('D-15', '2.2 ③ · 4.3', '이새 · 회사 API 범위 · 잘못 들어간 예외',
     '이새에프앤씨 8/30 예외가 회사 API 범위로 잘못 들어갔고(10/03 에이전트 범위로 다시 등록) 옛 줄이 아직 유효',
     '운영 조치', '이새가 회사 API 연동을 쓰지 않으면 옛 줄을 화면에서 회수(비활성 · 회수 기록이 남는다 · 사유 = 「범위 착오 · 에이전트 범위로 재등록」) · 쓰면 반박(같은 서버 주소 하나)',
     '닫힘(10-21 13:49 대표 회수 · 감사 기록 접속 예외 해제 · 사유는 대장 입력안 2차 [4] = D-35)', 'Q-38'),
    ('D-35', '2.2 ③', '예외 회수 사유 · 사유 빈칸',
     '국외 접속 예외 회수 버튼이 사유를 묻지 않아 감사 기록 사유가 null(10-21 이새 회수 실측 · GeoAccessTab.tsx:57 빈 본문)',
     '운영 조치 + 추가 과제', '이번 회수 사유는 운영 기록 대장에 남김(입력안 2차 [4]) · 회수 창 사유 입력칸은 실사 뒤 개선 과제(코드 · 이번 범위 아님)',
     '대장 기록 완료(10-21) · 추가 과제(회수 창 사유 칸)', 'Q-38'),
    ('D-16', '2.2 ⑤', '연동 경로 기록만 · 모든 접근 경로 · 출발지 차단',
     '한줄로 SDK · 싱크에이전트와 게이트웨이 API 연동은 미등록 출발지를 「기록만」 한다(문서 5절에 그대로 적었다 · 시행 스위치 없음 = H-24)',
     '운영 조치(게이트웨이 · 한줄로)', '게이트웨이 API = 10/03 뒤 7,944건 전부 사유 없음(목록 밖 0 · 출발지 1곳) → 코드가 정한 전환 조건 충족(client-auth.js 「목록 밖 0 확인 뒤 enforce」) → GW_API_ORIGIN_ALLOWLIST=enforce + 관리 API 재기동 + 전환 뒤 정상 고객 접수 계속 확인 · 9월 이후 REST 를 안 쓴 계정은 등록 출발지 밖이면 막힌다(정책대로) / 한줄로 = 회사 API 미등록 0 · 싱크에이전트 10월 미등록 2곳 → H-11b 보고 결정',
     '두 쪽 다 켤 조건 충족(한줄로 = 10/15 23:13 뒤 미등록 0 · H-11b) · 닫힘 10-21 = 한줄로 .env ORIGIN_ALLOWLIST_ENFORCE_FROM=2026-10-21 · 14:24:18 재기동 · 게이트웨이 drop-in zz-origin-allowlist.conf GW_API_ORIGIN_ALLOWLIST=enforce · 프로세스 확인 · 켠 뒤 게이트웨이 API 28건 ok · 막힘 0 · 한줄로 감지 · 차단 0', 'Q-41'),
    ('D-17', '2.3', '소명 없이 승인 · 47초 · 승인 근거',
     '10/20 보류 건(30만 원)을 고객사 소명 없이 47초 만에 승인 · 메모는 「무통장입금」뿐',
     '수정 + 반박', '수정(Harold 10-21 선택) = 소명 없는 보류 건 승인 때 「확인 근거」 입력 필수(돈 경로 · 설계안 → 동의 → Codex 적대) / 반박 = 입금자명에 등록 회사명이 들어 있음(데이터)',
     '배포 완료 10-21(Codex 적대 3R approve) · 실측 대기 · ⚠ 임상옥 = 등록 대표자인지', 'Q-45'),
    ('D-18', '2.3 · 개인정보', '입금자 이름 원문 · 조회 화면',
     '엑셀 H-12 둘째 조회가 처리 기록 상세를 통째로 찍어 입금자 이름이 원문으로 나왔다(실사 화면 노출)',
     '수정(조회문)', 'H-12 둘째 조회를 칸별로 바꾸고 입금자는 첫 글자만 · 저장된 원문은 반박(Q-48)',
     '생성기 반영 · 엑셀은 마지막에', 'Q-48'),
    ('D-19', '3.1 ③ · ④', '사설 원격 접속망 · Tailscale · 출장 되돌리기',
     '10/06 출장 때 두 서버에 사설 원격 접속망을 켜고 .62 방화벽에 줄을 더했다 · 귀국 뒤 되돌림과 운영 기록 대장 기록(켠 날 · 끈 날)이 확인되지 않는다',
     '운영 조치', '10-21 실측 = .62 · .65 둘 다 tailscaled active · enabled(귀국 뒤 되돌리지 않음) → 두 서버 끄기(접속 출발지 확인 뒤) · .62 trip-tailscale-ssh 줄 삭제 · 콘솔 접근 규칙 정리 → 대장 방화벽 정책 변경에 켠 날(10/06) · 끈 날(10/21) 사실대로 기록',
     '닫힘(10-21 대장 방화벽 정책 변경 4건 등록 10:29~10:32 · 대표 확인 10:37 · H-16)', 'Q-49'),
    ('D-20', '3.2', 'psypsy · 박성용 · 심사 계정',
     '8/27 발급 1분 뒤 중지된 계정 · 중지 사유가 「전송자격인증제 심사위한 계정생성」이라 보여주기용으로 읽힐 수 있다',
     '반박', 'Harold 10-21 = 잘못 생성해 정지 · 쓸 수 없던 계정(로그인 0 · OTP 없음 · 비밀번호 변경 필요 상태로 1분 뒤 중지) · 중지 사유 = 생성 경위',
     '닫힘', 'Q-52'),
    ('D-21', '공통', 'UTC 표기 · 9시간',
     '직원 계정 · 발신번호 표의 시각이 UTC 로 저장돼 그대로 찍으면 9시간 빠르게 보인다',
     '수정(조회문)', 'H-15 · H-8b 조회를 한국 시각으로 바꿈(AT TIME ZONE UTC) · 물으면 Q-55',
     '생성기 반영 · 엑셀은 마지막에', 'Q-55'),
    ('D-22', '3.4 ④', '계정 잠금 실례 없음 · 5회 실패',
     '인증번호 5회 실패 잠금이 실제로 일어난 적이 없어 동작 증거가 없다',
     '반박', 'Harold 10-21 = 시험 안 함 · 「실제 잠금은 아직 없었다 · 실패 9월 1건」',
     '답변 확정', 'Q-60'),
    ('D-23', '3.4 ② · ③', '본인확인 뒤 수동 변경 · gwchae',
     '10/13 ceo 가 본인확인을 마친 gwchae 의 인증번호 수신 번호를 바꿨다 · 문서 3.4 2절 「손으로 바꿀 수 없다」와 어긋날 수 있다 · 10/20 대상 빈칸 1건',
     '반박 + 수정 후보', '반박 = 10/13 전후 같음 · 본인확인 5계정 지금 번호 = 본인확인 번호(H-7b t 5/5) / 수정 후보 = 본인확인을 마친 계정은 관리자 수동 변경을 막거나 사유 필수(메모리 1007 「PUT /users/:id/mfa-phone 닫기」 남은 일)',
     '닫힘(Harold 10-21 = 수동 변경 경로 수정은 이번에 안 함)', 'Q-61'),
    ('D-24', '2.1 · 4.2', '계정 삭제 · 삭제 기록 · efolium',
     '계정 삭제 경로 두 곳(관리자 routes/admin.ts:500 · 고객사 관리자 routes/manage-users.ts:225)에 감사 기록이 없다 · 10/20 옛 efolium 계정 삭제가 기록 없이 사라졌다',
     '수정', 'D-03 과 같은 뿌리(삭제 경로가 감사 기록을 안 남김) → 2.1 · 4.2 에 걸리는 표(계정 · 발신번호 · 배정 · 수신거부 등)의 삭제 경로를 전수 grep 해 한 번에 감사 기록 연결',
     '배포 완료 10-21(Codex 적대 3R approve) · 실측 대기 · 10/20 경위 = 이폴리움 요청으로 고객사 분리(선불 금액 따로 관리 · Harold 10-21)', 'Q-62'),
    ('D-25', '5.1 · 4.2', '고객사 해지 · 사유 · 고지 결과',
     '고객사 해지 기록에 사유와 고지 결과가 없다(문서 4.1-4.3 3-1 · 1.3 6-4 · 6-5 는 남는다고 적었다)',
     '수정', '해지 때 사유를 필수로 받고 고지 결과(문자 고지 성공 · 실패)를 감사 기록에 남김 · 지난 2건은 반박(계약 종료 해지 · ⚠ Harold)',
     '배포 완료 10-21 · 실측 완료(13:39 테스트주식회사 해지 · 문자 수신 · H-18b)', 'Q-73'),
    ('D-26', '5.2', '금칙어 차단 규칙 0 · 탐지만',
     '지금 차단 모드 금칙어 규칙이 하나도 없다(불법스팸 8/27~ 탐지 · 시험 규칙도 탐지로 되돌림)',
     '운영 조치', '불법스팸 규칙을 화면에서 차단으로 전환(두 달 탐지 · 오탐 0 = 문서 5-1 전환 조건) · 전환 기록이 남는다',
     '닫힘(10-21 13:53:44 불법스팸 차단 전환 · H-19)', 'Q-75'),
    ('D-27', '5.2', 'KISA 차단정보 · 외부 출처',
     '문서 5-1 「KISA 등 외부 차단정보는 출처를 표시해 등록」 · 외부 출처 규칙 0개',
     '반박', 'KISA 차단정보는 링크가드 관제센터에 출처 KISA 로 적재돼 발송 엔진에서 차단 중(10-21 관제센터 화면 · 마지막 적재 08:47) · 한줄로 금칙어 화면은 자체 문구 규칙',
     '닫힘(Harold 10-21 화면)', 'Q-77'),
    ('D-28', '4.2 · 4.3', 'DB 오류 로그 · 컬럼 없음 · cdp_push_subscriptions · imc_template_status',
     '운영 DB 로그에 몇 분마다 「컬럼 없음」 오류(cdp_push_subscriptions.is_active · kakao_templates.imc_template_status)가 찍힌다',
     '수정', '설계안 S5 = push 조회 2곳 status=\'active\' 로(10-21 실측 활성 구독 0건 → 판정 그대로) · imc_template_status 는 컬럼 존재 캐시로 실패 SQL 을 안 보냄(동작 그대로)',
     '배포 완료 10-21 · 실측 완료(재기동 뒤 is_active 0 · imc 재기동 직후 1건) · 추가 과제 = 웹푸시 구독 첫 행이 생길 때 채널 판정 확인', 'Q-80'),
    ('D-34', '5.1 · 3.3', '회사 해지 두 경로 · 수정 화면 해지 · 마트테스트',
     '회사 수정(PUT /companies/:id)으로 상태를 해지로 바꾸면 계정 잠금 · 세션 끊기 · 해지 기록 · 고지가 없다 · 해지 회사 21곳 중 해지 기록 2곳 · 마트테스트 계정 2개 안 잠김',
     '수정 + 운영 조치 + 반박', '설계안 S3 = 해지 한 경로 CT(사유 필수 · 잠금 · 세션 · 고지 · 기록) · 수정 화면 해지 거부 / 운영 조치 = 마트테스트 계정 2개 잠금 · 대장 기록 / 반박 = 지난 해지 시점은 기록이 없다(사실대로)',
     '배포 완료 10-21 · 실측 완료(H-18b) · 마트테스트 사람 계정 testmart 10-21 13:59 잠금(account_restricted · 번호 없어 고지 false) · 남은 1개 = 연동 계정 system_sync(로그인 자체가 막힌 시스템 계정 · Q-04) · 대장 로그 점검 10-21 등록 · 테스트주식회사(실측) 해지 유지(Harold) · 지난 해지 시점= 기록 없음(사실대로)', 'Q-73'),
    ('D-29', '3.1 · 4.1 ⑤', '외부 열린 DB · pay-ingest-db · 23388',
     'pay-ingest-db 가 0.0.0.0:23388 로 열려 있어 「DB 내부 전용」 문장과 어긋나 보인다',
     '반박', '결제 통계 수신 전용 · 문서 3.1 4절에 별도 행으로 적었다 · 10-21 실측 = DOCKER-USER 23388 송신 서버 5곳만 허용 뒤 DROP',
     '닫힘', 'Q-81'),
    ('D-31', '4.2 · 4.1 ⑤', '시스템 저널 8월 15일 · 1년 보관 · 파일 개수 한도 · SystemMaxFiles',
     '한줄로 서버 시스템 저널이 08-15 부터만 남았다 · 9/12 설정(1year · 12G) 때는 06-06 ~ 09-12 가 있었다 · 저널 파일이 정확히 100개 = 파일 개수 기본 한도(SystemMaxFiles)에 걸려 오래된 파일이 지워졌다 · 지금도 새 파일마다 하나씩 지워진다',
     '운영 조치 + 반박', '운영 조치(지금) = journald.conf 백업 · SystemMaxFiles=2000 추가 · journald 재시작 · 운영 기록 대장(로그 점검)에 경위 기록 / 반박 = 지워진 것은 서버 시스템 저널 일부 · 이용자 로그 · 감사 기록은 DB 에 그대로(2026-02-11~) · 게이트웨이 서버도 같은 설정이라 파일 개수 확인',
     '10-21 두 서버 설정 반영(SystemMaxFiles=2000 · journald active · 백업 journald.conf.bak-1021 · 게이트웨이는 85개에서 미리 막음 · 엔진 active) · 대장 로그 점검 10-21 등록 · 닫힘', 'Q-84'),
    ('D-32', '4.3', '점검 조회 · 관리자 로그인 0 · ADMIN_LOGIN_MFA',
     '게이트웨이 월간 점검 조회(G-10)가 기록에 없는 이름(ADMIN_LOGIN)으로 세어 「관리자 로그인 0」으로 보였다',
     '수정(조회문)', 'G-10 을 ADMIN_LOGIN_MFA(OTP 까지 마친 로그인)로 고침 · 9월 점검 기록(대장)은 88로 맞게 적혀 있다',
     '닫힘(10-21 고친 조회 확인 = 9월 88 · 10월 49)', 'Q-92'),
    ('D-33', '2.2 ①', '사설 주소 허용 IP · eternal',
     '게이트웨이 허용 IP 에 사설 주소(10.0.131.103 · 10.0.12.170)가 섞여 있다',
     '반박', 'Harold 10-21 = 정리 안 함 · 바깥 접속은 공인 주소로 판정',
     '답변 확정', 'Q-89'),
    ('D-30', '4.1 ⑤', '백업 비밀번호 인자 · mysqldump 경고',
     '백업 기록에 mysqldump 「명령줄 비밀번호」 경고가 매일 찍힌다',
     '반박', 'Harold 10-21 = 이번에 안 바꿈 · 서버 안 관리자 실행 · 원격 접속 두 곳 · 백업 파일 공개키 암호화',
     '답변 확정', 'Q-82'),
]

HEAD_FILL = PatternFill('solid', fgColor='1F3A5F')
HEAD_FONT = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')
CMD_FONT = Font(name='Consolas', size=10)


def sheet(ws, headers, rows, widths, cmd_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical='center', wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        if cmd_col:
            row[cmd_col - 1].font = CMD_FONT
    ws.freeze_panes = 'A2'


MD_OUT = Path(__file__).resolve().parent.parent / '실사_당일_명령표.md'
SURAN_OUT = Path(__file__).resolve().parent.parent / '서수란_실사_질의응답_20261022.xlsx'

# ───────────────────────────── 서수란 팀장 묶음(명령 없음) ─────────────────────────────
# 지원팀장 등급이 여는 화면만(utils/admin-role.ts 권한분류표 · 10-21 코드 확인). 감사 로그 · 본인인증 현황 · 감시 기록 = 대표 전용.
SURAN_SCREENS = [
    ('2.1 계정', '고객 관리 → 고객사 관리 · 사용자 관리', '고객사 목록과 상태(활성 · 정지 · 해지) · 고객 계정 목록과 상태', ''),
    ('2.1 ⑤ 발신번호', '발송 관리 → 발신번호 관리', '발신번호 관리 · 등록 신청 관리(서류 · 심사 · 반려 사유) · 등록현황 관리(배정)', ''),
    ('2.2 국외 이용자', '보안 · 인증 → 국외 접근 통제', '국외 IP 대역 전면 차단 · 관리자 수동 승인(예외 IP · 사유 · 만료) · 국외 IP 탐지 · 차단 결과 로그', ''),
    ('2.3 부정사용', '요금/정산 → 충전 관리', '승인 대기(명의 확인 보류 사유 · 고객사 소명) · 잔액 변동 이력', '승인 · 반려 버튼은 누르지 않는다'),
    ('3.1 방화벽 · 4.3 로그 점검 · 3.3 권한 점검', '보안 · 인증 → 운영 기록 대장', '로그 점검 · 방화벽 정책 변경 · 접근권한 점검 기록(작성자 · 대표 확인)', ''),
    ('3.2 · 3.3 접근권한', '보안 · 인증 → 직원 계정·권한', '계정 목록 · 권한분류표 · 접근권한 변경 이력 대장', '지원팀장은 조회만 된다'),
    ('3.4 다중인증 · 4.1 로그인', '보안 · 인증 → 로그인 차단 관리', '로그인 차단 목록 · 수동 차단', '추가 · 해제 버튼은 누르지 않는다'),
    ('5.1 조치', '고객 관리 → 고객사 관리', '고객사 상태(정지 · 해지) · 해지 버튼(사유 선택)', '해지 버튼은 누르지 않는다'),
    ('5.2 금칙어', '보안 · 인증 → 금칙어 차단', '금칙어 · 악성 URL 자동 차단 · 차단정보 등록 · 목록 · 탐지 · 차단 결과 로그', ''),
    ('4.1 · 4.2 이용자 · 시스템 로그 · 3.4 본인인증 · 3.5 발신 인증', '(대표 전용) 감사 로그 · 본인인증 현황 · 감시 기록', '서팀장 화면에는 메뉴가 없다(권한분류표대로 정상)', '「대표님이 로그로 보여 드리겠습니다」'),
]


def suran_screen(item: str) -> str:
    i = item.replace(' ', '')
    if i.startswith('2.1⑤'):
        return '발신번호 관리'
    for prefix, screen in (('2.1', '고객사 관리 · 사용자 관리'), ('2.2', '국외 접근 통제'), ('2.3', '충전 관리'),
                           ('3.1', '운영 기록 대장(방화벽 정책 변경)'), ('3.2', '직원 계정·권한'), ('3.3', '직원 계정·권한 · 운영 기록 대장(접근권한 점검)'),
                           ('3.4', '로그인 차단 관리(본인인증 현황은 대표)'), ('3.5', '대표가 로그로'), ('4.1', '대표가 로그로'),
                           ('4.2', '대표가 로그로'), ('4.3', '운영 기록 대장(로그 점검)'), ('5.1', '고객사 관리(상태)'), ('5.2', '금칙어 차단')):
        if i.startswith(prefix):
            return screen
    return '대표가 로그로'


def suran_answer(a: str) -> str:
    # 확인 전(⚠) · 덧붙일 문장이 남은 답은 서팀장이 말하지 않는다
    if '⚠' in a or '+' in a or not a.strip().startswith('「'):
        return '(대표님이 답합니다)'
    return a.strip()


def write_suran():
    rows = [(no, item, kw, q, suran_answer(a), suran_screen(item), show) for no, item, kw, q, a, show, basis in QA]
    wb = Workbook()
    ws = wb.active
    ws.title = '물어보면 답'
    sheet(ws, ['번호', '심사 항목', '핵심 검색어(Ctrl+F)', '예상 질문', '답변(그대로 말한다)', '내가 띄울 화면', '대표님이 띄울 로그 번호'],
          rows, [7, 10, 28, 40, 80, 26, 16])
    sheet(wb.create_sheet('띄울 화면'), ['심사 항목', '메뉴', '화면에 나오는 것', '주의'], SURAN_SCREENS, [30, 36, 70, 30])
    wb.save(SURAN_OUT)
    told = sum(1 for r in rows if r[4] != '(대표님이 답합니다)')
    print(f'저장: {SURAN_OUT} · 질답 {len(rows)}건(서팀장이 말할 것 {told} · 대표 {len(rows) - told})')


def _cell(v) -> str:
    return str(v if v is not None else '').replace('|', '／').replace('\n', ' ')


def write_md():
    """AI 가 당일 이 파일 하나만 읽고 답하도록 — 위에 번호 → 줄 포인터 · 아래에 본문(명령 · 질답 · 결함)."""
    head = [
        '# 실사 당일 명령표 · 2026-10-22 (생성본 · 손으로 고치지 않는다)',
        '',
        '> 원본 = `evidence/make_inspection_commands_xlsx.py` · 다시 만들기 = `python evidence/make_inspection_commands_xlsx.py`(엑셀과 함께 나온다)',
        '',
        '## 0. 답하는 순서 (AI)',
        '1. Harold 가 친 요청의 심사 항목 · 낱말로 §1 찾기표에서 번호와 **줄**을 찾는다(낱말 = 검색어 칸).',
        '2. 그 줄만 읽고(Read offset = 줄 · limit 12) **실행 위치 + 명령 한 줄을 그대로** 준다. 명령은 고치지 않는다 · 한 번에 하나.',
        '3. 말로 답할 질문이면 §1-3 에서 질답 번호를 찾아 「Q-nn · 보여 줄 번호」만 한 줄(말은 서수란 팀장).',
        '4. 표에 없으면 「처음 쓰는 명령」이라고 밝히고 SELECT 만 · SCHEMA 정독 뒤. ⚠ 표시 문장은 말하지 않는다.',
        '5. 금지 = sudo · git · 원격 접속 명령 · .env 출력 · docker exec -it · 비밀번호가 찍히는 명령.',
        '',
        '답 모양: **실행 위치** 한 줄 → 명령 코드 블록 하나 → 「보는 것」 한 줄.',
        '',
    ]
    body = []
    anchors = {}

    def mark(key):
        anchors[key] = len(body)

    body += ['## 2. 명령', '']
    for no, item, kw, title, where, cmd, basis, memo in CMD_ROWS:
        mark(no)
        body += [f'### {no} · {title}', f'- 항목 {item} · 검색어 {kw}', f'- **실행 위치** {where}', '```bash', cmd, '```']
        if memo:
            body.append(f'- 보는 것 · 메모: {memo}')
        body += [f'- 근거: {basis}', '']
    body += ['## 3. 물어보면 답(서수란 팀장이 말한다)', '']
    for no, item, kw, q, a, show, basis in QA:
        mark(no)
        body += [f'### {no} · {item} · {kw}', f'- 질문: {q}', f'- 답변: {a}', f'- 보여 줄 번호: {show}', f'- 근거 · 주의: {basis}', '']
    body += ['## 4. 결함 대장', '', '| 번호 | 항목 | 결함 | 처리 | 상태 | 질답 |', '|---|---|---|---|---|---|']
    for no, item, kw, defect, how, why, state, qa in DEFECTS:
        body.append(f'| {no} | {_cell(item)} | {_cell(defect)} | {_cell(how)} | {_cell(state)} | {_cell(qa)} |')

    idx = ['## 1. 찾기표', '', '### 1-1. 심사 항목 → 명령 번호', '', '| 심사 항목 | 한줄로 | 게이트웨이 |', '|---|---|---|']
    idx += [f'| {_cell(a)} | {_cell(b)} | {_cell(c)} |' for a, b, c in FIND]
    idx += ['', '### 1-2. 명령 번호 → 줄', '', '| 번호 | 항목 | 어떤 로그 | 검색어 | 실행 위치 | 줄 |', '|---|---|---|---|---|---|']
    cmd_rows_at = len(idx)
    idx += [''] * len(CMD_ROWS)
    idx += ['', '### 1-3. 질답 → 줄', '', '| 번호 | 항목 | 검색어 | 보여 줄 번호 | 줄 |', '|---|---|---|---|---|']
    qa_rows_at = len(idx)
    idx += [''] * len(QA)
    idx += ['']
    base = len(head) + len(idx) + 1  # 본문 첫 줄의 줄 번호
    for i, (no, item, kw, title, where, cmd, basis, memo) in enumerate(CMD_ROWS):
        idx[cmd_rows_at + i] = f'| {no} | {_cell(item)} | {_cell(title)} | {_cell(kw)} | {_cell(where)} | {base + anchors[no]} |'
    for i, (no, item, kw, q, a, show, basis) in enumerate(QA):
        idx[qa_rows_at + i] = f'| {no} | {_cell(item)} | {_cell(kw)} | {_cell(show)} | {base + anchors[no]} |'

    lines = head + idx + body
    MD_OUT.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    # 자가 검증 — 포인터 줄이 실제로 그 번호의 제목 줄인가
    got = MD_OUT.read_text(encoding='utf-8').split('\n')
    for key, at in anchors.items():
        assert got[base + at - 1].startswith(f'### {key} · '), (key, base + at, got[base + at - 1][:40])
    print(f'저장: {MD_OUT} · 포인터 {len(anchors)}개 확인')


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = '명령집'
    sheet(ws, ['번호', '심사 항목', '핵심 검색어', '어떤 로그', '실행 위치', '명령 (칸 클릭 → Ctrl+C → 터미널에 붙여넣기 · 붙이면 바로 실행된다)', '근거', '메모'],
          CMD_ROWS, [8, 12, 30, 32, 22, 110, 16, 40], cmd_col=6)
    sheet(wb.create_sheet('오늘 리허설'), ['번호', '실행 위치', '무엇', '명령', '보는 것'], REHEARSAL, [6, 24, 34, 110, 50], cmd_col=4)
    sheet(wb.create_sheet('항목별 찾기'), ['심사 항목', '한줄로', '게이트웨이'], FIND, [44, 60, 50])
    sheet(wb.create_sheet('물어보면 답'), ['번호', '심사 항목', '핵심 검색어', '예상 질문', '답변(그대로 말할 문장)', '보여 줄 번호', '근거 · 주의(⚠ = 확인 전 · Harold 답)'],
          QA, [7, 10, 26, 40, 80, 16, 50])
    sheet(wb.create_sheet('결함 대장'), ['번호', '심사 항목', '핵심 검색어', '결함', '처리(수정 · 운영 조치 · 반박)', '처방 또는 반박 근거', '상태', '관련 질답'],
          DEFECTS, [7, 10, 24, 50, 16, 70, 24, 14])
    wb.save(OUT)

    # 자가 검증 — 명령 칸은 전부 한 줄 · 수식으로 읽히는 값 없음
    from openpyxl import load_workbook
    chk = load_workbook(OUT)
    n = 0
    for name, col in (('명령집', 6), ('오늘 리허설', 4)):
        for row in chk[name].iter_rows(min_row=2):
            v = row[col - 1].value or ''
            assert '\n' not in v and not str(v).startswith('='), f'{name} {row[0].value}'
            n += 1
    print(f'저장: {OUT} · 명령 칸 {n}개 · 전부 한 줄')
    write_md()
    write_suran()


if __name__ == '__main__':
    sys.exit(main())
